"""Finish wording, refresh embedded figures, and align public package documentation."""
from pathlib import Path
from docx import Document
from PIL import Image
from docx.shared import Inches
import copy

R=Path(__file__).resolve().parents[1]
h=(R/'code/condense_results_v107.py').read_text()
exec(h[h.index('def replace_span('):h.index('doc = Document(MANUSCRIPT)')])
p=R/'manuscript/MANUSCRIPT_NEUROCHEMISTRY_INTERNATIONAL_FINAL.docx'; d=Document(p)
replacements={
 'was unchanged in ALS spinal cord':'showed no statistically significant difference in ALS spinal cord',
 'TRPC1 remained unchanged in the well-powered cervical and lumbar spinal-cord comparisons':'TRPC1 showed no statistically significant difference in cervical or lumbar spinal cord',
 'within groups it followed neuronal content rather than the junction.':'within groups its association with the junction was attenuated after adjustment for a neuronal marker.',
 'so it does not appear to be only a difference in cell composition;':'although marker adjustment cannot exclude residual effects of cell composition;',
 'because its mean Ct was essentially unchanged':'because its mean Ct was similar',
 'SUPPA2 gave a more conservative result on the same libraries.':'SUPPA2 returned fewer calls on the same libraries.',
 'SUPPA2 builds its null from the between-replicate ΔPSI distribution, which makes it conservative in a three-versus-three design, so the disagreement bounds how much weight any single event call can carry rather than showing that one tool is wrong.':'SUPPA2 uses between-replicate ΔPSI variation to construct its null. The limited agreement indicates method sensitivity; this comparison does not identify which method is better calibrated.',
}
for pp in d.paragraphs:
    for a,b in replacements.items():
        if a in pp.text: replace_span(pp,a,b)
# Source figures have been regenerated. Use the preceding image paragraph of each
# caption rather than page numbers, which can change after editing.
def image_before(doc,caption,png,width=None):
    idx=next(i for i,p in enumerate(doc.paragraphs) if p.text.startswith(caption))
    candidates=[p for p in doc.paragraphs[max(0,idx-3):idx] if p._p.xpath('.//a:blip')]
    assert len(candidates)==1,(caption,len(candidates))
    pp=candidates[0]; blip=pp._p.xpath('.//a:blip')[0]
    rid=blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
    pp.part.related_parts[rid]._blob=png.read_bytes()
    exts=pp._p.xpath('.//wp:extent')+pp._p.xpath('.//a:ext')
    cx=int(Inches(width)) if width else int(exts[0].get('cx'))
    w,h=Image.open(png).size; cy=int(cx*h/w)
    for e in exts: e.set('cx',str(cx)); e.set('cy',str(cy))
image_before(d,'Figure 1.',R/'figures/main/Figure1_functional_consequences.png',6.05)
image_before(d,'Figure 6.',R/'figures/main/Figure6_working_model.png',6.05)
d.save(p)
p=R/'supplementary/SUPPLEMENTARY_MATERIAL.docx'; s=Document(p)
image_before(s,'Supplementary Figure S5.',R/'figures/supplementary/Supplementary_Figure_S5_NMD_interaction.png')
for pp in s.paragraphs:
    if pp.text.startswith('S1.'):
        replace_span(pp,pp.text,'S1. Laboratory source data: raw Ct and relative expression from four technical RT-qPCR measurements per group, Fura-2 amplitudes from three wells per group on one culture plate, WST-1 values from four wells of one experiment, primers and thermal profile. All laboratory summaries are descriptive.')
s.save(p)

p=R/'DATA_AVAILABILITY.md'; p.write_text(p.read_text().replace('v1.0.5','v1.0.6'))
p=R/'README.md'; s=p.read_text()
s=s.replace('Five\nanalyses were revised','Several\nanalyses were revised')
s=s.replace('`S10_NMD_interaction_SOCE_panel.csv`, `S10b_NMD_panel_level_tests.csv`','`S10_NMD_interaction_SOCE_panel.csv`, `S10b_NMD_panel_descriptive_summary.csv`')
s=s.replace('S10b_NMD_panel_level_tests.csv','S10b_NMD_panel_descriptive_summary.csv')
s=s.replace('source_data/nmd_panel_t4.csv','source_data/nmd_descriptive_all_genes.csv.gz')
s=s.replace('the RT-qPCR measurement','the descriptive RT-qPCR measurement')
s=s.replace('For the current figure set, the laboratory source values are read by\n`code/fig_main_lab.py` and `code/build_source_data.py`;','For the current figure set, the laboratory source values are read by\n`code/fig_lab_descriptive_v118.py`; `code/relabel_s1_qpcr_v118.mjs` maintains the\nsource workbook’s technical-replicate labels and descriptive summaries;')
s=s.replace('under which no gene passes genome-wide FDR.','whose inferential p values have been withdrawn. Current NMD tables contain descriptive estimates only.')
s += '''
## September 29 corrections and current rebuild order

Release v1.0.6 includes the scientific audit corrections described in
`CORRECTIONS_2026-09-29.md`. RT-qPCR has four technical measurements per group,
Fura-2 three wells per group on one plate, and the available WST-1 values four
wells from one experiment. None of these laboratory comparisons is assigned an
inferential p value. Public RNA-seq experiments retain their own biological designs.

Current analysis/figure entrypoints (with the workstation input paths configured):

1. `python code/audit_tables_v118.py`: S13, descriptive S10/S10b, full NMD source table,
   and all nine STIM1 transcript-test eligibility records.
2. `python code/nygc_donor_sensitivity_v118.py`: S18d donor sensitivity; also reproduces
   S18/S18b/S18c through the shared normalization module.
3. `node code/relabel_s1_qpcr_v118.mjs`: preserve S1 measurements and update experimental units.
4. `python code/fig_lab_descriptive_v118.py` and `python code/fig_nmd_descriptive_v118.py`:
   current Figure 1 and Supplementary Figure S5.
5. `python code/export_manuscript_md.py`, `python code/qa_check_v4.py` and
   `python code/qa_audit_v118.py`: export and verify the edited manuscript and supporting files.

`recalc_nmd_shared_control.R` now delegates to the descriptive builder. Old editing
scripts and the historical nine-figure sources record earlier revisions; do not use
them to overwrite the current reviewed DOCX. S18d corrects across regions within
group/gene/model, whereas Table 5 corrects across genes within region.
'''
p.write_text(s)
p=R/'CHANGES_FROM_THESIS.md'; s=p.read_text(); a=s.index('## 2. NMD interaction'); b=s.index('## 3.',a)
s=s[:a]+'''## 2. NMD interaction: descriptive contrasts with shared controls

Both the original eight-value test and the later four-condition t/sign tests reuse
the same baseline libraries. Averaging within conditions did not remove covariance
or include uncertainty in the shared baseline difference. Their inferential p/q
values are withdrawn. Current S10 and S10b report mean, range, condition estimates
and panel medians without inferential tests. CBARP remains a positive descriptive
pattern (+1.52 log2 across four contrasts), not an established NMD target.
The 16 cryptic-splicing reference genes are not a validated NMD-positive panel;
their summary cannot establish assay failure. The full descriptive gene table is
`source_data/nmd_descriptive_all_genes.csv.gz`.

'''+s[b:]
s+='''
## 6. Experimental units and September 29 scientific audit

The author confirmed that the Fura-2 observations are three wells from one plate,
and the four RT-qPCR observations are technical repeats of the same biological
sample per group. All laboratory comparisons are now descriptive; qPCR and Fura-2
p values, biological confidence intervals and significance stars are omitted.
WST-1 remains four wells from one of three experiments; the other experiments are
not available. See `CORRECTIONS_2026-09-29.md` for the S13 coordinate/chromosome fix,
STIM1 missing tests, DE threshold clarification and NYGC donor sensitivity.
'''; p.write_text(s)
print('Final wording, figure embeds and public documentation aligned')
