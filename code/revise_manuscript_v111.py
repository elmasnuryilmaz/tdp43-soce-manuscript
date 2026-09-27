"""Resolve the internal contradictions found in the 27 September structural scan.

A1  the editorial sentence "This is the central observation of the study" is removed; a referee
    had asked for it and the data make the point without it
A2  Section 3.8 no longer ends on a conclusion its own numbers contradict: two nominal p values
    among 24 nested combinations are what chance gives
A3  the Discussion no longer says the matched null removed the panel enrichment, because in the
    motor neurons it did not
A4  the heading of Section 3.4 drops "concentrate", which implied an enrichment that Section 3.8
    denies and that this section never tested
A5  the Introduction previews the polyadenylation and NMD analyses, so Sections 3.6 and 3.7 no
    longer arrive unannounced
"""
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "manuscript/MANUSCRIPT_NEUROCHEMISTRY_INTERNATIONAL_FINAL.docx"
src = (ROOT / "code/condense_results_v107.py").read_text()
exec(src[src.index("def replace_span("):src.index("doc = Document(MANUSCRIPT)")])

doc = Document(MANUSCRIPT)
p = list(doc.paragraphs)
assert p[86].text == "3.4 Robust splicing changes concentrate in SOCE regulators"

# A1
replace_span(p[69], " This is the central observation of the study.", "")

# A4
replace_span(p[86], "3.4 Robust splicing changes concentrate in SOCE regulators",
             "3.4 Which splicing changes in SOCE genes survive the robustness checks")

# A2
replace_span(p[119], "We therefore regard enrichment of splicing changes in Ca²⁺ homeostasis genes "
                     "as a tentative, motor-neuron-associated observation rather than a general "
                     "property of TDP-43 depletion.",
             "Two nominal results among 24 dataset–panel combinations, drawn from panels that are "
             "nested by construction, is what chance gives. Once the number of combinations is "
             "taken into account there is no evidence that Ca²⁺ homeostasis genes are enriched for "
             "splicing changes, and the apparent enrichment in the primary model is explained by "
             "the length and expression properties of those genes. We report the two motor-neuron "
             "panels as the only combinations that survived matching, not as an established "
             "enrichment.")

# A3
replace_span(p[141], "Read-support and matched-background analyses were decisive here: they removed "
                     "isolated events and apparent panel enrichment while retaining CBARP and a "
                     "small number of model-specific candidates.",
             "Read-support filtering and the matched-background null did most of the work here: "
             "together they removed isolated events and, in the primary model and the iPSC "
             "colonies, the apparent panel enrichment, while retaining CBARP and a small number of "
             "model-specific candidates. In the iPSC-derived motor neurons two panels survived "
             "matching, which is within what 24 nested combinations produce by chance.")

# A5
replace_span(p[11], "We then use public RNA-seq datasets to examine calcium-regulatory transcript "
                    "abundance and RNA processing, applying read-support and robustness checks to "
                    "the splicing results.",
             "We then use public RNA-seq datasets to examine calcium-regulatory transcript "
             "abundance and the three routes by which TDP-43 loss is known to change RNA "
             "processing: annotated and cryptic splicing, alternative polyadenylation, which "
             "generates the truncated STMN2 transcript, and coupling to nonsense-mediated decay. "
             "Read-support and robustness checks are applied throughout to the splicing results.")

doc.save(MANUSCRIPT)
print("revised:", MANUSCRIPT.name)
