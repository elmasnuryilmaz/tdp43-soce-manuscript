"""Apply the v1.0.7 referee-style evaluation to the manuscript, supplement and highlights (release v1.0.8).

Changes by theme (see CORRECTIONS_2026-09-29_v1.0.8.md):
  * claim narrowed to the evidence: the Fura-2 result is a single-plate discovery observation that
    is not paired with the RNA data; the RNA candidates are the main content
  * 'high-confidence' -> 'stringent-filter'; 'SOCE machinery / regulators / core entry components'
    -> the fixed gene lists they stood for
  * the detection-power simulation is described with its assumptions
  * svaseq numbers, STIM1 transcript identifiers and the APA details move to the supplement;
    the NMD and FRASER results share one short paragraph with the APA screen (Section 3.6)
  * Discussion reordered: RNA findings first, single-plate hypotheses condensed, then tissue
  * figure captions rewritten for the redrawn Figures 2, 4, 5 and 6; alt text added

The script is deliberately strict: every replacement must find its text exactly once, and the
number of Zotero citation fields must be unchanged by the edits that only move or reword text.

Run once from the repository root:  /usr/bin/python3 code/revise_manuscript_v122.py [--out DIR]
"""
import copy
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx_edit_tools as T  # noqa: E402

R = Path(__file__).resolve().parents[1]
OUT = R
if "--out" in sys.argv:
    OUT = Path(sys.argv[sys.argv.index("--out") + 1])
MAIN_IN = R / "manuscript/MANUSCRIPT_NEUROCHEMISTRY_INTERNATIONAL_FINAL.docx"
SUPP_IN = R / "supplementary/SUPPLEMENTARY_MATERIAL.docx"
HL_IN = R / "highlights_Neurochemistry_International.docx"
NEW_TITLE = ("Calcium-regulatory RNA candidates after TDP-43 knockdown in SH-SY5Y cells: reanalysis of public "
             "RNA-seq data and a single-plate Ca²⁺-readdition observation")
OLD_TITLE = ("TDP-43 knockdown is associated with a lower Ca²⁺-readdition response and altered "
             "calcium-regulatory RNA profiles in SH-SY5Y cells")

# =========================================================================== main manuscript
doc = Document(str(MAIN_IN))
P = list(doc.paragraphs)
fields_before = T.field_count(doc)


def expect(i, opening):
    assert P[i].text.startswith(opening), (i, P[i].text[:60], opening)
    return P[i]


def edit(i, old, new, opening=None):
    p = expect(i, opening) if opening else P[i]
    T.edit_text(p, old, new)
    return p


def norm(p):
    """Re-apply gene italics after a text edit (keeps citations; not for code-styled runs)."""
    assert not p._p.xpath('.//w:rStyle[@w:val="VerbatimChar"]') and not p._p.xpath(".//w:vertAlign")
    T.rebuild(p, T.to_markup(p))


# capture the caption run formatting (Times New Roman, 9.5 pt) used by Figures 2, 3 and 5
_cap_run = next(a["el"] for a in T.atoms(P[69]) if a["kind"] == "run" and not a["text"].startswith("Figure"))
CAP = T._clean_rpr(_cap_run)


def caption(i, opening, markup):
    p = expect(i, opening)
    T.rebuild(p, markup, template_rpr=CAP)
    return p


# ------------------------------------------------------------------ title, abstract, introduction
edit(0, OLD_TITLE, NEW_TITLE, "TDP-43 knockdown is associated")

expect(5, "TDP-43 loss is implicated")
T.rebuild(P[5], (
    "TDP-43 loss is implicated in amyotrophic lateral sclerosis (ALS), but its relationship to store-operated "
    "Ca²⁺ entry (SOCE) is unclear. After shRNA-mediated *TARDBP* depletion in SH-SY5Y cells, *TRPC1*, *STIM1*, "
    "*ORAI1* and *ATP2A3* mRNA measurements were 1.7- to 3.2-fold higher (four biological RT-qPCR replicates per "
    "group). In one Fura-2 culture plate (three wells per group; descriptive), the Ca²⁺-readdition amplitude was "
    "84% lower in knockdown wells, ER Ca²⁺ release 33% lower and the readdition-to-release ratio 77% lower; this "
    "observation motivated the RNA analyses but was not paired with the RT-qPCR or RNA-seq samples. The 48-h "
    "WST-1 signal was 38.5% lower (four wells, one experiment; descriptive). We reanalysed six public "
    "TDP-43-depletion RNA-seq comparisons (38 libraries) for calcium-regulatory expression and splicing, screened "
    "eleven comparisons for unannotated junction changes, and examined alternative polyadenylation and "
    "nonsense-mediated decay descriptively. *CBARP* was the most reproducible splicing candidate across species, "
    "whereas *STIM2*, *STIMATE* and *ORAI3* events were supported in the primary RNA-seq model. No confirmed "
    "RNA-processing event in the core SOCE-pathway genes explained the Fura-2 difference. In ALS tissue, *TRPC1* "
    "and *SARAF* increased and *CBARP* decreased in six brain regions, and *SARAF* and *CBARP* changed in the same "
    "direction in cervical and lumbar cord, but these tissue associations could not be attributed to TDP-43 loss. "
    "The RNA analyses identify candidate calcium-regulatory changes without establishing a causal link to the "
    "Fura-2 response, which requires independent biological replication."))

expect(11, "We first measure SOCE")
T.rebuild(P[11], (
    "We first measured the Ca²⁺-readdition response of a store-depletion–readdition protocol, and target mRNAs, in "
    "separate SH-SY5Y cultures after shRNA-mediated TDP-43 depletion; the Fura-2 observation comes from a single "
    "culture plate and served to motivate the RNA analyses, not to establish a phenotype. We then used public "
    "RNA-seq datasets to examine calcium-regulatory transcript abundance and the three routes by which TDP-43 "
    "loss is known to change RNA processing: annotated and cryptic splicing, alternative polyadenylation, which "
    "generates the truncated STMN2 transcript, and coupling to nonsense-mediated decay. Read-support and "
    "robustness checks are applied throughout to the splicing results. Finally, we asked how the candidate "
    "transcripts behave in ALS brain and selected neurological comparison cohorts."))

# ------------------------------------------------------------------ methods
# 2.2: the svaseq numbers move to Results 3.2 and Supplementary Results 5
p19 = expect(19, "Gene counts were generated")
tail_start = p19.text.index(", which estimated two surrogate variables")
T.edit_text(p19, p19.text[tail_start:],
            ", which estimated two surrogate variables in the primary SH-SY5Y comparison; the refit with these "
            "variables is reported in Section 3.2 and Supplementary Results 5.")
# 2.2: every DESeq2 statistic reported in the article comes from the Salmon counts of the primary comparison;
# the featureCounts run of the same libraries is described as the cross-check it is
edit(19, "Gene counts were generated with featureCounts v2.1.1 ",
     "Differential expression was analysed on gene-level counts. For the primary SH-SY5Y comparison (GSE296712; "
     "0 versus 75 ng/mL doxycycline, three libraries per group), every DESeq2 statistic reported here was "
     "computed on Salmon estimated counts summed to genes with the complete transcript-to-gene map (Section 2.4). "
     "A count of the same six libraries with featureCounts v2.1.1 ")
edit(19, " at meta-feature level against the same annotation (",
     " at meta-feature level against the GENCODE v47 basic annotation (")
edit(19, " for paired-end libraries; multi-mapping and multi-overlapping reads not counted) and modelled with DESeq2 v1.42.1 ",
     " for paired-end libraries; multi-mapping and multi-overlapping reads not counted) gave similar log2 fold "
     "changes (Pearson r = 0.85 across the 20,833 genes counted in both). Counts were modelled with DESeq2 v1.42.1 ")
edit(19, " within each dataset, with Benjamini–Hochberg correction ", ", with Benjamini–Hochberg correction ")

edit(22, "the twelve core entry components", "the twelve core SOCE-pathway genes", "Because rMATS applies FDR")

expect(23, "Detection power was estimated")
T.rebuild(P[23], (
    "For Supplementary Figure S1, detection power was simulated for a 3 + 3 design at fixed depths of 10, 20, 50 "
    "and 100 informative reads per sample: replicate PSI was drawn from a normal distribution around 0.5 ± ΔPSI/2 "
    "(standard deviation 0.05, limited to 0.01–0.99), reads were sampled binomially, groups were compared with a "
    "two-sample t test at nominal p < 0.05, and each condition was simulated 2,000 times. The simulation does not "
    "reproduce the rMATS model, its FDR correction or its |ΔPSI| threshold."))

expect(29, "The calling threshold")
T.rebuild(P[29], (
    "The calling threshold was calibrated against a control-versus-control null: in datasets with four control "
    "replicates, the controls were split 2 + 2, the analysis repeated unchanged, and every call counted as a false "
    "positive. The permissive definition required an unannotated junction, ΔPSI ≥ 0.05, q < 0.05, control PSI ≤ "
    "0.05 and a bootstrap lower bound above zero (Supplementary Table S14). Because this definition produced about "
    "as many split-control calls as real calls, or more (Section 3.5), we adopted an effect-based definition: a "
    "**stringent-filter unannotated splicing candidate** requires a novel splice site, ΔPSI ≥ 0.20, q < 0.05, a "
    "bootstrap lower bound above 0.05, at least 20 reads in the knockdown group, and non-zero counts in every "
    "knockdown replicate. The name describes the filter and does not imply validated specificity. A requirement of "
    "zero counts in controls was deliberately not imposed, because it removes genuinely cryptic exons with basal "
    "leakiness, *STMN2* among them. Positive-control recovery, not call count, is used throughout as the measure "
    "of whether an analysis detects known targets; because the sixteen literature controls are human cryptic "
    "events, this check is available only for the human comparisons."))

expect(32, "FRASER v1.14.1")
T.rebuild(P[32], (
    "FRASER v1.14.1 [[cite:Mertes]] was run on the nine-sample SH-SY5Y doxycycline series (0, 25 and 75 ng/mL), "
    "counting split and non-split reads from the alignments and modelling ψ5, ψ3 and splicing efficiency θ. The "
    "method treats aberrant splicing as a rare deviation from a cohort norm, whereas two thirds of this cohort are "
    "depleted samples; the run is therefore treated as a limit of the approach (Supplementary Results 1)."))

expect(39, "Four cumulative gene panels")
T.rebuild(P[39], (
    "Four cumulative gene panels were assembled from curated SOCE/TRP components (n = 51), channels and transport "
    "systems (n = 117), curated Ca²⁺-handling genes (n = 258) and an expanded Ca²⁺-associated set from KEGG and "
    "Gene Ontology (n = 732; Supplementary Table S2). Symbols were harmonised to HGNC/MGI and orthology verified "
    "via Ensembl Compara. Where results are summarised for the nineteen-gene SOCE-associated set, this refers to a "
    "fixed list: *STIM1*, *STIM2*, *ORAI1–3*, *TRPC1*, *SARAF*, *STIMATE*, *CBARP*, *CRACR2A*, *CRACR2B*, "
    "*SELENOK*, *ATP2A1–3*, *MCU*, *MCUB*, *MICU1* and *MICU2*; the core SOCE/TRP panel is Tier 1. The list "
    "combines core STIM–ORAI components with store-refilling, mitochondrial-uptake and voltage-gated-channel-"
    "associated genes (for example *ATP2A1–3*, *MCU* and *CBARP*). *CBARP* was included in this fixed list for the subsequent cross-dataset event ranking. The "
    "twelve genes examined event by event in Section 2.3 (*STIM1*, *STIM2*, *STIMATE*, *SARAF*, *CRACR2A*, "
    "*CRACR2B*, *ORAI1–3*, *TRPC1*, *ATP2A2* and *ATP2A3*) are a subset of these nineteen and are called the core "
    "SOCE-pathway genes."))

edit(42, "TRPC, SOCE regulators, mitochondrial", "TRPC, Ca²⁺-entry regulators, mitochondrial", "Transcripts per million")
norm(P[42])
edit(44, "The junction-level analysis of Section 3.9", "The junction-level analysis of Section 3.8", "ALS post-mortem")

# ------------------------------------------------------------------ results 3.1
edit(61, "3.1 TDP-43 knockdown is associated with a smaller Ca²⁺-readdition amplitude despite increased SOCE-related mRNAs",
     "3.1 SOCE-associated mRNAs are higher, and the Ca²⁺-readdition amplitude is lower in one Fura-2 plate, after "
     "TDP-43 knockdown", "3.1 TDP-43 knockdown")

caption(65, "Figure 1.", (
    "**Figure 1.** Laboratory measurements in SH-SY5Y cells. (A) *TARDBP* mRNA in untransduced, non-targeting shRNA "
    "and shTDP-43 samples. (B) Relative *TRPC1*, *STIM1*, *ORAI1* and *ATP2A3* mRNA. RT-qPCR groups comprise four "
    "biological replicates; panels A and B use separate RNA sets. (C) WST-1 signal at 48 h, four wells from one "
    "experiment (descriptive). Bars are means with SEM across biological replicates (A, B) or wells (C); dots are "
    "individual measurements. RT-qPCR: two-sided Welch tests on ΔCt with Holm correction (Methods 2.16); adjusted "
    "p: *TARDBP*, 8.3 × 10⁻¹¹ versus non-targeting shRNA and 3.2 × 10⁻¹⁰ versus untransduced cells; *TRPC1*, "
    "0.0040; *STIM1*, 0.0040; *ORAI1*, 0.0026; *ATP2A3*, 6.8 × 10⁻⁵."))

edit(66, "The Fura-2 Ca²⁺-readdition amplitude was markedly lower in shTDP-43 cells (Figure 2).",
     "In the single Fura-2 plate, the Ca²⁺-readdition amplitude was markedly lower in shTDP-43 wells (Figure 2).",
     "The Fura-2")
edit(66, "similar baselines before readdition (Figure 2E)", "similar baselines before readdition (Figure 2A)")

expect(67, "The measured mRNAs")
T.rebuild(P[67], (
    "The RT-qPCR mRNA measurements and the one-plate Fura-2 response therefore differed in direction; because they "
    "come from different cultures and time points (Methods 2.13 and 2.14), they are not paired and do not show "
    "opposing changes within the same cells."))

caption(69, "Figure 2.", (
    "**Figure 2.** Fura-2 measurements in SH-SY5Y cells after TDP-43 knockdown (one culture plate). (A) One "
    "representative recording per group on common axes, aligned to the steepest point of the Ca²⁺-readdition rise "
    "(dashed line); no value is smoothed or rescaled. The original recordings, with the CPA and CaCl₂ additions, "
    "are shown at their own axis ranges in Supplementary Figure S9. (B) ER Ca²⁺ release after CPA (10 µM), (C) "
    "Ca²⁺-readdition amplitude after CaCl₂ (nominally 1.5 mM) and (D) the readdition-to-release ratio of each well. "
    "Bars are means ± well-to-well SEM and dots are the three wells per group; both phases of a well come from one "
    "recording. The comparisons are descriptive and no inferential test is shown."))

# ------------------------------------------------------------------ results 3.2
edit(70, "3.2 Transcript-family abundance contextualises the discordance",
     "3.2 Transcript-family composition in the public SH-SY5Y RNA-seq model", "3.2 Transcript-family")
edit(71, "The independent public SH-SY5Y RNA-seq comparison provides transcript-level context for this laboratory observation (Figure 3; Table 1).",
     "The public SH-SY5Y RNA-seq comparison, which is separate from our cultures, provides transcript-level context "
     "for the laboratory measurements (Figure 3; Table 1).", "The independent public")
edit(73, "SOCE-regulator pool", "Ca²⁺-entry-regulator pool", "In control cells")
edit(73, "The STIM (+48%) and SOCE-regulator (+30%) pools", "The STIM (+48%) and Ca²⁺-entry-regulator (+30%) pools")
norm(P[73])

expect(74, "For STIMATE, adjusted TPM")
T.rebuild(P[74], (
    "For *STIMATE* and *MCUR1*, adjusted TPM and the DESeq2 estimate differ in sign (+9.9% versus log2FC −0.050, "
    "p_adj = 0.72; −7.8% versus +0.044, p_adj = 0.62). Both changes are small, reflect different estimators and do "
    "not support a directional abundance effect."))
T.insert_after(P[74], (
    "Adjustment for two surrogate variables estimated with svaseq (Methods 2.2) kept the direction of all 850 genes "
    "that met both differential-expression thresholds in the two models but reduced the number meeting them from "
    "1,694 to 1,067. *STIM1*, *TRPC1*, *ORAI3*, *SARAF* and *CBARP* retained q < 0.05; *ORAI1*, *ATP2A3*, *ATP2A2* "
    "and *STIM2* did not, so the RNA-seq support for the *ORAI1* and *ATP2A3* increases is model dependent, whereas "
    "the RT-qPCR measurements of Section 3.1 support their direction (Supplementary Results 5)."), like=P[74])

expect(75, "Higher ORAI3 and SARAF transcripts")
T.rebuild(P[75], (
    "Higher *ORAI3* and *SARAF* transcripts are candidate routes to reduced entry, whereas lower *ORAI2* would on the "
    "same evidence favour entry: __ORAI2__ and __ORAI3__ form heteromers with __ORAI1__ and restrain SOCE "
    "[[cite:Vaeth]], and __SARAF__ facilitates slow Ca²⁺-dependent inactivation [[cite:Palty]]. The "
    "mitochondrial-uptake transcripts also changed, but MCU, MCUB and MICU proteins have distinct, context-dependent "
    "roles [[cite:Lambert]], and neither summed transcript changes nor the *STIM1*:*ORAI1* RNA ratio measure "
    "protein composition at ER–plasma-membrane junctions [[cite:Hoover]]. These changes provide hypotheses for the "
    "lower Ca²⁺-readdition signal observed in one laboratory plate, not an explanation of it."))
expect(76, "Together, the changes")
T.delete(P[76])

caption(78, "Figure 3.", (
    "**Figure 3.** Calcium-regulatory transcript profile in the public SH-SY5Y RNA-seq comparison (RNA level only). "
    "(A) Adjusted gene-level TPM in control and TDP-43-depleted libraries (three per group), on a logarithmic "
    "scale; lines link the two group estimates of a gene. (B) DESeq2 gene-level log2 fold changes; filled markers "
    "indicate Benjamini–Hochberg q < 0.05 (adjusted p values in Table 1). Bold labels mark the four RT-qPCR targets. "
    "These libraries are distinct from the cultures of Figures 1 and 2. TPM and fold changes describe transcript "
    "abundance, not protein or function, and can differ in sign when a change is small (Section 3.2)."))

# ------------------------------------------------------------------ results 3.3
edit(79, "but most individual events are underpowered", "but many individual events rest on limited read support",
     "3.3 TDP-43 depletion")

expect(82, "Simulation over the observed coverage")
T.rebuild(P[82], (
    "An example simulation (Methods 2.3; Supplementary Figure S1) illustrates the consequence. For a 3 + 3 design "
    "with replicate PSI varying with a standard deviation of 0.05 around 0.5, 10 informative reads per sample, "
    "close to the lower quartile of the nominally significant events (9.7 reads per sample; median 21.3), gave a "
    "power of 0.08 to detect a true ΔPSI of 0.10 at nominal p < 0.05, and 100 reads per sample gave 0.28. A power "
    "of 0.80 was not reached at 10 or 20 reads per sample for any simulated ΔPSI up to 0.30, and at 50 and 100 "
    "reads per sample it required a ΔPSI between 0.20 and 0.30. Nominally significant events at low coverage are "
    "therefore expected to be inflated in effect size, and individual events at the |ΔPSI| ≥ 0.10 threshold should "
    "not be interpreted without read-level support. The simulation is not an estimate of rMATS power."))

expect(83, "We applied this reasoning")
T.rebuild(P[83], (
    "As an example of why read-level verification matters, a *TRPC1* skipped-exon event (chr3:142,792,824–"
    "142,792,967; ΔPSI = +0.108, FDR = 0.0459) passed the conventional thresholds but rested on seven skipping "
    "reads across six libraries, with the skipping form absent in four of six samples. Its bootstrap 95% "
    "confidence interval spanned zero (−0.093 to +0.522), removing one control replicate reversed the sign of ΔPSI "
    "(−0.074), the event did not survive coverage pre-filtering, and LeafCutter did not call it. We therefore do "
    "not report it as a finding (Supplementary Figure S2)."))

# ------------------------------------------------------------------ results 3.4
edit(84, "3.4 Which splicing changes in SOCE genes survive the robustness checks",
     "3.4 Splicing changes in core SOCE-pathway genes that pass the robustness checks", "3.4 Which splicing")

expect(85, "Applying the robustness criteria")
T.rebuild(P[85], (
    "Applying the robustness criteria to the twelve core SOCE-pathway genes examined event by event (Methods 2.3) "
    "identified three events with adequate coverage and bootstrap intervals excluding zero: *STIMATE*, *ORAI3* and "
    "*STIM2* (Table 3; Figure 4A). The *STIM2* event is a 119-bp exon and is distinct from the 24-nucleotide "
    "*STIM2*.1 exon analysed in Section 3.5 (Figure 4B)."))

expect(86, "In the primary SH-SY5Y model these were")
T.rebuild(P[86], (
    "In the primary SH-SY5Y model these were *STIMATE* (ΔPSI = +0.244; FDR = 9.0 × 10⁻⁷; 95% CI +0.095 to "
    "+0.368), *ORAI3* (−0.269; 1.0 × 10⁻²; −0.404 to −0.107) and *STIM2* (−0.120; < 1 × 10⁻¹⁶; −0.165 to "
    "−0.064). A *STIM1* event (chr11:4,088,702–4,088,738; 37 bp, frame-disrupting) reached ΔPSI = +0.145 (FDR = "
    "4.0 × 10⁻⁴) with a confidence interval marginally including zero, and was positive in all three human "
    "datasets."))
T.insert_after(P[86], (
    "Two separate isoform-level analyses addressed *STIM1* and are not a single confirmation chain. In the "
    "DRIMSeq–stageR analysis, *STIM1* passed the gene-level screening stage (q = 5.49 × 10⁻⁷), but none of seven "
    "evaluable transcripts passed stageR confirmation (all confirmation-stage adjusted p values = 1.0); two further "
    "transcripts had missing DRIMSeq p values and could not be evaluated. In the separate "
    "IsoformSwitchAnalyzeR/DEXSeq analysis, the two isoforms with significant usage changes were not predicted to "
    "carry a premature termination codon (PTC), whereas the two predicted PTC isoforms did not change (q = 0.33 "
    "and 0.77). No confirmed PTC-associated isoform switch was therefore identified; transcript identifiers and "
    "test eligibility are given in Supplementary Results 6."), like=P[86])

edit(87, "the largest expression change of any SOCE regulator in SH-SY5Y",
     "the largest expression change of any Ca²⁺-entry regulator of Table 1 in SH-SY5Y", "CBARP, encoding")
norm(P[87])

expect(88, "LeafCutter, which does not use")
T.rebuild(P[88], (
    "LeafCutter, which does not use the reference annotation to define events, also called *CBARP* (p_adj = 0.0133) "
    "and *STIM2* (p_adj = 0.0288) in the same SH-SY5Y libraries. The coverage-filtered rMATS analysis, the "
    "bootstrap interval estimates and LeafCutter, all applied to the same libraries, therefore point to the same "
    "candidates; this is agreement between analysis methods, not independent replication."))

caption(90, "Figure 4.", (
    "**Figure 4.** Splicing evidence in SOCE-related genes. (A) Per-library PSI for the four skipped-exon events of "
    "Table 3 (public SH-SY5Y comparison, three libraries per group); each title gives the exon length, lines show "
    "group means, and the text gives the rMATS ΔPSI with its replicate-level bootstrap 95% interval. The *STIM1* "
    "interval includes zero. PSI is the share of inclusion-junction reads among inclusion and skipping reads "
    "(schematic). (B) A different *STIM2* exon: the 24-nucleotide SOAR exon that converts *STIM2* into *STIM2*.1, "
    "as an rMATS event in six datasets; squares are sized by inverse-variance weight, bars are 95% intervals and "
    "diamonds are fixed-effect pooled estimates for all six and for the three human datasets (Supplementary Table "
    "S15). (C) *CBARP*: share of exon-4 donor reads joined to the alternative 3′ splice site at chr19:1,235,342 "
    "rather than to exon 5, per library (Methods 2.5; the full locus is in Supplementary Figure S3)."))

# ------------------------------------------------------------------ results 3.5
edit(91, "3.5 Annotation-free analysis recovers cryptic exons genome-wide but finds no high-confidence event in the SOCE machinery of the primary model",
     "3.5 Annotation-free junction analysis recovers known cryptic exons but finds no stringent-filter unannotated "
     "change in the core SOCE-pathway genes of the primary model", "3.5 Annotation-free")
edit(92, "on a second, independently produced junction set (Methods 2.5)",
     "on a second junction set produced with a different extractor (Methods 2.5)", "rMATS and LeafCutter both")
edit(93, "and 117 high-confidence unannotated splicing changes in 87 genes were called",
     "and 117 stringent-filter unannotated splicing candidates in 87 genes were called", "The approach recovered")
edit(93, "gave 165 high-confidence events in 113 genes", "gave 165 stringent-filter candidates in 113 genes")
norm(P[93])
expect(94, "Specificity was tested")
T.rebuild(P[94], (
    "Specificity was tested with FUS and TAF15 knockdown in the same iPSC-derived motor neurons, at the same depth "
    "and with the same design (Supplementary Figure S4C). Under the permissive definition the three comparisons "
    "produced call counts of the same order (141 for TDP-43, 124 for FUS and 126 for TAF15), but only TDP-43 "
    "knockdown recovered positive controls, two of the sixteen (*STMN2* and *KALRN*); two controls against none is "
    "not a significant difference (Fisher’s exact test p = 0.48), and this dataset is too shallow for the "
    "stringent-filter threshold to recover a control in any of the three comparisons (18, 26 and 23 events, none "
    "of them a control gene). The matched comparison therefore does not establish specificity and remains "
    "descriptive. Deeper TDP-43 comparisons recovered thirteen of sixteen positive controls in SH-SY5Y and fifteen "
    "in iPSC colonies, supporting sensitivity to known targets in those models."))

expect(95, "The null test also sets")
T.rebuild(P[95], (
    "The null test sets the limits of interpretation. Under the permissive rule, the null-to-real call ratios were "
    "0.98 in iPSC colonies and 2.01 in K562 total RNA (Supplementary Table S14). The stringent-filter null could be "
    "run in the three datasets with four control replicates: for every real call the split-control null produced "
    "0.64 calls in iPSC colonies, 2.17 in K562 total RNA and 0.83 in mouse striatum (Supplementary Table S14). "
    "These ratios are not calibrated false-discovery rates, because sample size and the number of tests differ "
    "between the two analyses, but they show that call counts are not interpretable on their own, and that the "
    "stringent-filter list is a set of candidates whose sensitivity has been checked against known targets, not a "
    "set with established specificity. In the remaining eight comparisons the controls could not be split, and "
    "positive-control recovery is the only available check. Across the eleven comparisons the burden of "
    "stringent-filter candidates ranged from 12 to 477 events, and of the three datasets with a null, the one with "
    "the largest burden also had the largest null in absolute number."))

expect(96, "Applied to the Ca²⁺ panels")
T.rebuild(P[96], (
    "Applied to the Ca²⁺ panels, the analysis returned an almost complete negative (Supplementary Table S7), which "
    "is informative where the analysis demonstrably worked: in SH-SY5Y at both doses and in K562 poly(A)+ mRNA, "
    "where positive controls were recovered at the stringent-filter threshold, and less strongly in the TDP-43 "
    "knockdown of iPSC-derived motor neurons, where two were recovered under the permissive definition only. In "
    "K562 total RNA no positive control was recovered, and the three mouse comparisons have no conserved control, "
    "so the absence of calls there carries little information. The only stringent-filter unannotated change anywhere in the "
    "nineteen-gene SOCE-associated set was the *CBARP* junction in iPSC colonies (the core SOCE/TRP panel also "
    "carried one in *TRPM3* in that comparison); *STIM1*, *STIM2*, *ORAI1*–3, *TRPC1*, *SARAF*, *STIMATE* and the "
    "SERCA and mitochondrial-uptake genes carried none in any comparison. That single call comes from the dataset "
    "whose null test produced 0.64 calls for every real call, so it is not evidence on its own, although the "
    "*CBARP* junction is corroborated below. This finding does not exclude lower-abundance or context-specific "
    "events, but it provides no support for a cryptic-splicing switch as the explanation for the within-plate "
    "Fura-2 difference (permissive-tier and annotated-site events: Supplementary Table S7 and Figure S4B)."))

expect(97, "What the annotation-free analysis")
T.rebuild(P[97], (
    "Separately from the genome-wide candidate list, the *CBARP* junction was evaluated directly at junction level "
    "in the two human models. In SH-SY5Y it showed the largest local change of any gene of the nineteen-gene "
    "SOCE-associated set, in the junction from exon 4 to the alternative 3′ splice site (ΔPSI +0.74, 95% CI +0.51 "
    "to +0.83, q = 2.7 × 10⁻¹²; +0.76, q = 5.6 × 10⁻¹⁷ in the regtools set), and one arm of the event uses a "
    "splice site absent from the annotation (Supplementary Figure S3). This is the third analysis of the same "
    "libraries, after rMATS with --novelSS and LeafCutter, to identify this locus; isoform-level testing did not "
    "detect a *CBARP* isoform switch (IsoformSwitchAnalyzeR gene-level q = 0.12; DRIMSeq q = 0.34). The same "
    "unannotated splice site was also found in the iPSC colonies (ΔPSI +0.33, q = 4.3 × 10⁻¹⁶, 201 versus 137 "
    "reads), where it met the cryptic criteria outright. *ORAI2*, the most abundant ORAI-family gene, showed +0.16 "
    "(q = 1.9 × 10⁻³) in the mapping-quality-filtered set and +0.15 (q = 0.087) in the unfiltered set, so the two "
    "extraction methods provide unequal support."))

expect(98, "One specific mechanism")
T.rebuild(P[98], (
    "One specific mechanism could be tested directly. Inclusion of a 24-nucleotide exon in the STIM–ORAI activating "
    "region (SOAR) converts *STIM2* into *STIM2*.1 (STIM2β), an isoform that inhibits SOCE [[cite:Miederer]], so "
    "increased inclusion upon TDP-43 loss would be a simple route to reduced entry. Measured as an rMATS event, "
    "inclusion was higher in knockdown in five of six datasets, but the fixed-effect meta-analysis estimate was "
    "negligible (pooled ΔPSI +0.0013, 95% CI −0.022 to +0.024; p = 0.914; Figure 4B; Supplementary Table S15). The "
    "three mouse datasets carry 86% of the weight, but the human datasets alone give the same answer (+0.031, 95% "
    "CI −0.030 to +0.091; p = 0.32), and the estimates are not heterogeneous (Cochran’s Q = 4.84, 5 df, p = 0.44; "
    "I² = 0%). The pooled estimate is a fixed-effect summary across species and cell models, not an equivalence "
    "test, and it does not exclude an effect in a single model. At junction level, both flanks of the exon "
    "(chr4:27,007,983–27,008,006 in humans) were considered: sixteen flanking-junction records were measurable "
    "across seven TDP-43 comparisons and the FUS/TAF15 controls, and direction was positive in six TDP-43 "
    "comparisons and negative in C2C12 (Supplementary Table S13). The downstream estimate was +0.149 in "
    "SH-SY5Y at 75 ng/mL (q = 0.033, 20 versus 3 reads), +0.116 at 25 ng/mL (q = 0.14) and +0.020 in iPSC "
    "colonies (q = 0.82; 441 versus 281 reads), and it exceeds the rMATS estimate for the same exon (+0.149 versus "
    "+0.059) because rMATS combines both flanks and normalises by effective length. These junction summaries do not "
    "alter the separately computed rMATS meta-analysis, which did not detect a shared shift."))

# ------------------------------------------------------------------ results 3.6 (APA, NMD and FRASER together)
for i, start in ((99, "An outlier-based method"), (100, "3.6 Descriptive NMD"), (101, "The TDP-43 × NMD-inhibition")):
    expect(i, start)
    T.delete(P[i])
edit(102, "3.7 A coverage-based APA screen yields candidate gradients but no confirmed APA event in the SOCE machinery",
     "3.6 APA, NMD and outlier screens yield candidates but no confirmed event in the core SOCE-pathway genes",
     "3.7 A coverage-based")
expect(103, "Cryptic polyadenylation accounts")
T.rebuild(P[103], (
    "Cryptic polyadenylation accounts for a substantial part of TDP-43-dependent RNA processing "
    "[[cite:Bryce-Smith]]. The truncated *STMN2* transcript is generated this way, but APA had not been assessed "
    "for the Ca²⁺ gene set. We measured it from coverage, as the 5′ share of the two ends of each intron (an "
    "intronic polyadenylation index) and as the distal share of each terminal exon (a 3′UTR usage index), across "
    "the 696 qualifying genes of the expanded Ca²⁺ panel and the cryptic positive controls (Methods 2.7; "
    "Supplementary Figure S6A; Supplementary Table S11)."))

expect(104, "Excluding CIGAR reference skips")
T.rebuild(P[104], (
    "In the primary SH-SY5Y model the positive control responded only moderately: the *STMN2* intronic index "
    "shifted by +0.145 over all six libraries (bootstrap 95% CI +0.085 to +0.205; Supplementary Results 3). Among "
    "the depth-qualified units of the nineteen-gene SOCE-associated set, two intronic units showed moderate "
    "decreases (*STIM1* intron 17, Δ = −0.173, 95% CI −0.233 to −0.123; *STIM2* intron 13, Δ = −0.156, −0.239 to "
    "−0.037). Neither is independent of the splicing results: the 5′ window of the *STIM2* unit contains the "
    "alternatively spliced 119-bp *STIM2* exon of Table 3, and the *STIM1* unit is the intron immediately upstream "
    "of the alternatively included *STIM1* exon of Table 3. Both are candidate coverage gradients rather than "
    "localised poly(A) sites. No index change of 0.05 or more was found for *ORAI1*–3, *TRPC1*, *SARAF*, "
    "*STIMATE*, *CBARP* or the SERCA genes, and three terminal-exon units outside this set also excluded zero "
    "(*ITPR3*, −0.152; *MICU3*, +0.097; *ITPR1*, −0.055; Supplementary Results 3). Because the control responded "
    "only moderately, absence of a signal is not evidence that a gene is spared."))

expect(105, "Among the depth-qualified units")
T.rebuild(P[105], (
    "The analysis was repeated in the iPSC-derived motor neurons, where 59 units in 29 genes passed the depth "
    "filter and the positive control behaved as intended: the index of *STMN2* intron 2, which contains cryptic "
    "exon 2a, rose from 0.570 in controls to 0.819 in knockdown (Δ = +0.249, interval +0.208 to +0.290). Genes of "
    "the SOCE-associated set showed only small shifts, the largest being *ATP2A2* intron 3 (−0.164) and, below "
    "0.10, *SARAF* intron 5 (+0.097), *TRPC1* intron 1 (+0.083) and the *STIM2* terminal exon (−0.074); with n = 2 "
    "per group the enumerated bootstrap intervals are coarse (Supplementary Results 3)."))
for i, start in ((106, "The analysis therefore identifies"), (107, "We therefore repeated")):
    expect(i, start)
    T.delete(P[i])

expect(108, "The mouse datasets yielded")
T.rebuild(P[108], (
    "In the two mouse lines (74 qualifying units in C2C12 and 131 in NSC34) no positive control is available "
    "because the *STMN2* cryptic polyadenylation site is absent from the mouse gene [[cite:Melamed]], so negative "
    "results have limited interpretability. No unit of the SOCE-associated set exceeded |Δ| = 0.30. Units labelled "
    "intron 5 at the human *SARAF* locus (+0.097 in iPSC-derived motor neurons) and the mouse *Saraf* locus "
    "(+0.204 in NSC34) were matched by ordinal label, not sequence alignment, and remain gene-level candidates "
    "that do not demonstrate recurrence of the same RNA event (Supplementary Results 3)."))

expect(109, "Taken together, these analyses")
T.rebuild(P[109], (
    "The TDP-43 × NMD-inhibition experiment (GSE307054) was examined descriptively because its four contrasts share "
    "baseline libraries and TDP-43 status is confounded with sequencing batch (Methods 2.8; Supplementary Results "
    "2; Supplementary Figure S5). *CBARP* had the largest mean interaction among the genes in Supplementary Table "
    "S10 (+1.52 log₂; range across conditions +1.19 to +2.23), a candidate pattern and not evidence that *CBARP* is "
    "an NMD target; the cryptic-splicing reference genes were not selected as validated NMD-rescue controls, so "
    "their panel summary does not establish assay sensitivity or make a negative result conclusive (Supplementary "
    "Table S10b). The outlier-based FRASER analysis returned no genome-wide significant outlier, a design for "
    "which the method is poorly suited (Supplementary Results 1)."))
T.insert_after(P[109], (
    "Taken together, these screens support altered calcium-related RNA profiles but do not establish that any "
    "SOCE-pathway gene is a direct target of cryptic splicing, APA or NMD-coupled degradation. The APA and NMD "
    "analyses are hypothesis-generating: they measure coverage gradients rather than poly(A) sites, and their "
    "contrasts share controls while batch is confounded with TDP-43 status."), like=P[109])

# ------------------------------------------------------------------ results 3.7 and 3.8 (renumbered)
edit(110, "3.8 Matched analyses", "3.7 Matched analyses", "3.8 Matched")
edit(115, "3.9 TRPC1, SARAF and CBARP differ", "3.8 TRPC1, SARAF and CBARP differ", "3.9 TRPC1")

caption(124, "Figure 5.", (
    "**Figure 5.** ALS tissue expression of *TRPC1*, *SARAF* and *CBARP* by region. Cliff’s δ compares ALS with "
    "non-neurological controls within each NYGC region (case minus control) and is not adjusted for cell "
    "composition, age, RNA integrity or post-mortem interval; * marks Benjamini–Hochberg q < 0.05, corrected "
    "within each region across the tested genes. Row labels give the numbers of ALS cases and controls (exact q "
    "values in Table 5); the rule separates brain from spinal cord. The regions come from one consortium cohort "
    "and are analysed separately; they are not independent cohorts, and the associations do not establish that "
    "TDP-43 loss caused them."))

expect(121, "TRPC1 was also examined")
T.rebuild(P[121], (
    "*TRPC1* was also examined in four independent datasets covering Alzheimer’s disease, Parkinson’s disease and "
    "multiple sclerosis (Supplementary Figure S7). It was decreased in Alzheimer’s disease (δ = −0.447; q = 1.6 × "
    "10⁻⁷, with Braak-stage correlation ρ = −0.183, p = 1.8 × 10⁻³) and in Parkinson’s disease (δ = −0.677; q = 1.8 × "
    "10⁻⁴). In multiple sclerosis, *TRPC1* was decreased at donor level across all sampled lesion types (δ = "
    "−0.840; q = 0.038) and in lesions at sample level (δ = −0.594; q = 1.3 × 10⁻⁴), with the same direction in the "
    "second multiple-sclerosis cohort of five donors per group (five regions pooled, δ = −0.226; q = 0.49). Because "
    "*TRPC1* contributes to SOCE in oligodendrocyte precursor cells [[cite:Paez]], the multiple-sclerosis decrease "
    "could reflect demyelination. Two sensitivity analyses assessed that explanation: the decrease was present in "
    "normal-appearing white matter, where no statistically significant myelin-marker changes were detected "
    "(*TRPC1* δ = −0.482 at sample level, p = 0.005, q = 0.10; −0.771 at donor level, uncorrected p = 0.030), and "
    "regression on *MBP*, *PLP1* and *GFAP* attenuated it, leaving it significant at sample level but not at donor "
    "level (p = 0.055). These analyses do not exclude a contribution from tissue composition (Supplementary "
    "Results 4; Supplementary Table S16)."))

expect(129, "Across the cohort")
T.rebuild(P[129], (
    "Across groups and regions, *TRPC1* expression therefore did not track the cryptic *STMN2* indicator of TDP-43 "
    "loss of function: it rose in ALS brain regions, in which the junction exceeded 1% of reads in at most 14% of "
    "samples, showed no statistically significant difference in ALS spinal cord (41–67%) and fell in the cortex of "
    "the comparison group (50–63%), and within groups its association with the junction was attenuated after "
    "adjustment for a neuronal marker. This proxy analysis is correlational and does not test whether TDP-43 loss "
    "changes *TRPC1* expression."))

# ------------------------------------------------------------------ discussion
expect(131, "In one Fura-2 culture plate")
T.rebuild(P[131], (
    "This study asked which calcium-regulatory RNA changes accompany TDP-43 depletion and whether an RNA-processing "
    "event in the core SOCE-pathway genes could explain the lower Ca²⁺-readdition response seen in one Fura-2 "
    "culture plate. The most robust RNA findings were the recurrent *CBARP* splicing change, the shifts in ORAI- "
    "and STIM-family transcript composition in the public SH-SY5Y model, and the region-dependent *TRPC1*, *SARAF* "
    "and *CBARP* differences in ALS tissue. No confirmed RNA-processing event in the core SOCE-pathway genes "
    "explained the Fura-2 difference, which comes from three wells of one plate, was not measured in the RT-qPCR "
    "or RNA-seq cultures and requires independent replication. The RNA data are therefore candidates for matched, "
    "independently replicated experiments, not a mechanism for that observation."))

expect(137, "RNA-processing specificity.")
T.rebuild(P[137], (
    "**RNA-processing findings depend on model and method.** The annotation-free analysis recovered established "
    "TDP-43-dependent cryptic targets, including the *STMN2* and *UNC13A* programmes described across neuronal "
    "systems [[cite:Brown et al., 2022]], which supports detection of known targets in the deeper TDP-43 "
    "comparisons. Two controls versus none after FUS or TAF15 knockdown did not establish specificity, and the null "
    "comparisons show that the stringent-filter list is a set of candidates. The core SOCE-pathway genes carried no "
    "equivalent cryptic-splicing programme: the APA screen showed only coverage gradients and the NMD estimates "
    "were descriptive. Canonical loss of RNA repression is therefore reproducible across suitable neuronal models, "
    "whereas the calcium-regulatory analyses identify a smaller set of context-dependent processing events and "
    "broader changes in transcript-family composition."))

expect(138, "CBARP, TRPC1 and junctional regulators.")
T.rebuild(P[138], (
    "***CBARP* and the other splicing candidates.** *CBARP* is the strongest recurrent candidate: it was affected "
    "in five of six datasets across two species, identified by complementary analyses of the same libraries, "
    "reproduced at junction level in a second human dataset, and showed the largest expression change among the "
    "Ca²⁺-entry regulators of Table 1 (Figure 4C; Supplementary Figure S3). Its product BARP suppresses "
    "voltage-gated Ca²⁺ channel activity and Ca²⁺-evoked exocytosis [[cite:Béguin]]; this function concerns "
    "voltage-gated rather than store-operated channels, but it suggests that TDP-43 loss may remodel more than one "
    "route of calcium entry and secretion. The apparent *TRPC1* splicing event failed replicate-level robustness "
    "checks, so any *TRPC1* contribution, which has been linked to ER stress and AKT/mTOR signalling in neuronal "
    "cells [[cite:Selvaraj]], is more likely to involve abundance, assembly or cellular context than a splice "
    "switch. The pooled effect of the inhibitory *STIM2*.1 exon [[cite:Miederer]] was negligible (Supplementary "
    "Table S15; not a test of equivalence), so a shared *STIM2*.1 switch appears unlikely to explain the Fura-2 "
    "difference, whereas the *STIMATE* event, since __STIMATE__ promotes __STIM1__ activation at "
    "ER–plasma-membrane junctions [[cite:Jing]], and the *ORAI3* event remain plausible candidates. Read-support "
    "filtering and the matched-background null accounted for most of the reduction in candidate events, including "
    "the apparent panel enrichment in the primary model and the iPSC colonies."), extra=(P[139],))
expect(139, "STIM2, STIMATE and ORAI3.")
T.delete(P[139], citations_reused=True)
for q in (P[138], P[137]):
    T.move_after(q, P[131])

expect(132, "Interpretation of the Fura-2 response.")
T.rebuild(P[132], (
    "**Possible contributors to the single-plate Fura-2 difference.** Within the single plate, mean ER release was "
    "33% lower and mean readdition 84% lower, and readdition remained 77% lower when each well was normalised to "
    "its own release. This ratio is descriptive and cannot distinguish altered ER store depletion from changes in "
    "STIM activation, Ca²⁺ influx or clearance, and the response requires replication in independent cultures. The "
    "public transcript data suggest layers that could act together but were not tested. *ATP2A2*, which dominates "
    "the SERCA family, decreased modestly and could reduce ER refilling. *ORAI3* rose from 8% to 30% of ORAI "
    "transcripts, *ORAI2* fell from 77% to 53% and *SARAF* increased; because __SARAF__ promotes slow "
    "Ca²⁺-dependent inactivation and __ORAI2__ and __ORAI3__ can modify __ORAI1__-mediated kinetics and restrain "
    "entry in heteromeric channels [[cite:Palty et al., 2012]], the *ORAI3* and *SARAF* increases are plausible "
    "contributors to a smaller readdition amplitude, whereas the *ORAI2* decrease predicts the opposite. __ORAI3__ "
    "is not intrinsically inhibitory, however: in astrocytes __STIM1__ with __ORAI1__ and __ORAI3__ mediates most "
    "SOCE [[cite:Kwon]], and __SARAF__ also regulates entry in SH-SY5Y cells [[cite:Albarran]]. Channel output "
    "depends on protein abundance, assembly and localisation rather than summed RNA, and even the "
    "__STIM1__:__ORAI1__ protein ratio changes CRAC-channel trapping and gating [[cite:Hoover]]; a transcript "
    "rise could also be compensatory. A larger PMCA transcript pool (+35.2%, with *ATP2B2* and *ATP2B3* rising "
    "from low baseline) could accelerate extrusion, and because SERCA was inhibited by CPA during readdition, "
    "faster ER re-uptake cannot explain the peak without evidence of residual SERCA activity or inhibitor "
    "washout."), extra=(P[133], P[134]))
T.delete(P[133], citations_reused=True)
T.delete(P[134], citations_reused=True)

expect(135, "Comparison with ALS calcium phenotypes.")
T.rebuild(P[135], (
    "**Comparison with ALS calcium models and cell state.** SOCE dysregulation differs in direction across ALS "
    "models: SOD1(G93A) astrocytes show enhanced store-operated entry and Ca²⁺-dependent exocytosis "
    "[[cite:Kawamata]], with reduced SERCA abundance and lower resting ER Ca²⁺ in spinal-cord astrocytes of the "
    "same model [[cite:Norante]]. If the lower readdition signal of one SH-SY5Y plate were replicated, genetic "
    "lesion, cell identity, differentiation state and compensatory timing would be candidate reasons for a "
    "different direction. The 38.5% lower WST-1 signal (four wells, one experiment) is consistent with reported "
    "reductions in metabolic activity, growth and ATP production after *TARDBP* silencing "
    "[[cite:Ceron-Codorniu]], but it cannot separate fewer cells from lower activity per cell; the approximately "
    "25% lower mitochondrial-uptake transcript pool (mainly *MCU*, *MICU2* and *MCUB*) cannot be translated into "
    "flux [[cite:Lambert]]. Reduced ATP supply could impair SERCA-dependent store refilling, and altered "
    "mitochondrial uptake could change local Ca²⁺ clearance; direct measurements would be needed to test either."),
    extra=(P[136],))
T.delete(P[136], citations_reused=True)

expect(140, "Relevance to disease tissue.")
edit(140, "matching the direction observed in the public SH-SY5Y model.",
     "matching the direction observed in the public SH-SY5Y model. These are regional disease associations in bulk "
     "post-mortem tissue; age, RNA integrity and post-mortem interval were not modelled, and the diagnoses of the "
     "comparison group are not public.")
edit(140, "it did not, however, follow the cryptic STMN2 indicator of TDP-43 loss of function,",
     "in these data it did not, however, track the cryptic STMN2 indicator of TDP-43 loss of function,")
norm(P[140])

expect(141, "Integrated mechanism and testable predictions.")
T.rebuild(P[141], (
    "**Working hypotheses and discriminating experiments.** Earlier work linked TDP-43 to Ca²⁺ signalling through "
    "ER–mitochondrial contacts: disease-associated TDP-43 disrupts VAPB–PTPIP51 tethering [[cite:Stoica]], and "
    "increasing VAPB or PTPIP51 restores mutant-TDP-43-associated Ca²⁺ transfer and synaptic function "
    "[[cite:Markovinovic]]. The working model in Figure 6 proposes hypotheses about the plasma-membrane "
    "replenishment phase, none of which was tested here: TDP-43 loss might alter bioenergetics, ER refilling, "
    "channel composition and feedback control, and recurrent *CBARP* processing may add a parallel effect on "
    "voltage-gated entry and exocytosis. Each hypothesis makes a discriminating prediction. *ATP2A2* or "
    "bioenergetic perturbations could be tested with direct ER Ca²⁺ measurements and refilling after CPA washout; "
    "*SARAF* or ORAI perturbations for effects on the readdition response in independently replicated cultures; "
    "and an isoform-specific *CBARP* perturbation for a larger effect on depolarisation-evoked Ca²⁺ entry or "
    "secretion than on CPA-triggered SOCE."))

caption(143, "Figure 6.", (
    "**Figure 6.** Working model linking laboratory measurements to candidate hypotheses. Blue boxes and solid "
    "arrows: measurements after *TARDBP* knockdown, namely *TARDBP* mRNA and higher *TRPC1*, *STIM1*, *ORAI1* and "
    "*ATP2A3* mRNA (RT-qPCR, four biological replicates per group, day-5 RNA) and, in three wells per group of one "
    "culture plate 72 h after transduction, lower mean ER-release and Ca²⁺-readdition amplitudes (descriptive). "
    "The RT-qPCR and Fura-2 measurements come from different cultures and time points and are not paired. Orange "
    "boxes: candidate mechanisms suggested by public RNA-seq data; dashed arrows: hypotheses for perturbation, not "
    "demonstrated links. *CBARP* sits on a parallel voltage-gated branch that was not measured and has no arrow to "
    "the readdition response."))

# ------------------------------------------------------------------ limitations, conclusion
expect(145, "Lower TARDBP mRNA was measured")
T.rebuild(P[145], (
    "*TARDBP* depletion was shown at the mRNA level only, with a single TDP-43-targeting shRNA and no rescue "
    "experiment. The Fura-2 comparison rests on three wells per group from one culture plate prepared 72 h after "
    "transduction while puromycin selection was in progress; it does not estimate between-experiment variability, "
    "was not paired with the RNA samples and was not characterised pharmacologically, for example with a "
    "store-operated channel blocker such as BTP2, Synta66 or Gd³⁺, so the readdition signal is defined by the "
    "depletion–readdition protocol and not by pharmacology. Cuvette ratios report net cytosolic accumulation, and "
    "neither cell number nor maximum response was measured; the WST-1 signal (48 h, four wells of one experiment) "
    "is not a cell count, so a smaller or less healthy cell population could contribute to the smaller Fura-2 "
    "signal, although the ratiometric readout is in principle independent of cell number. RT-qPCR used four "
    "biological replicates per group and supported the direction of the four target-gene changes in an "
    "independent RNA sample set, but it was normalised to a single reference gene, *GAPDH*, whose mean Ct was "
    "similar between groups and which rose in the public inducible comparison (log2FC +0.61; *TBP* and *B2M* were "
    "stable), so the increases may be underestimated; amplification efficiencies were not determined, and *TARDBP* "
    "mRNA was not re-measured in the RNA set used for the target genes. The laboratory assay used undifferentiated "
    "SH-SY5Y cells, whereas the public study used an inducible model and the motor-neuron datasets had no matched "
    "Ca²⁺ measurements."))

expect(148, "In one SH-SY5Y culture-plate experiment")
T.rebuild(P[148], (
    "Reanalysis of public TDP-43-depletion RNA-seq data and laboratory RT-qPCR identified calcium-regulatory "
    "candidates: a recurrent *CBARP* splicing change, shifts in ORAI- and STIM-family transcript composition, and "
    "region-dependent *TRPC1*, *SARAF* and *CBARP* differences in ALS tissue. A lower Fura-2 Ca²⁺-readdition "
    "response was seen in knockdown wells of one culture plate, but it was neither replicated nor paired with the "
    "RNA measurements, and no confirmed RNA-processing event in the core SOCE-pathway genes explained it. "
    "Coverage-based APA and descriptive NMD summaries provide no confirmatory evidence that these genes are direct "
    "targets of those pathways. *CBARP*, *SARAF* and *TRPC1* are therefore priorities for protein-level and "
    "independently replicated functional follow-up."))

# ------------------------------------------------------------------ data availability
expect(154, "Data availability.")
T.retarget_hyperlinks(doc, P[154], "v1.0.7", "v1.0.8")
edit(154, "are provided in Supplementary Table S1.",
     "are provided in Supplementary Table S1. The values behind Supplementary Figure S1 and the surrogate-variable "
     "comparison of Section 3.2 are provided as source_data/power_simulation_S1.csv and "
     "source_data/svaseq_sensitivity_SHSY5Y.csv, and the featureCounts run compared in Section 2.2 as "
     "source_data/DESeq2_ctrl_vs_75_featureCounts.csv.")

# Table 1 note: the DESeq2 columns come from the Salmon counts
edit(256, "log2FC and p_adj are from DESeq2 gene-level counts;",
     "log2FC and p_adj are from DESeq2 on gene-level Salmon counts;", "Note. TPM, transcripts per million")

# ------------------------------------------------------------------ supplementary inventory in the main document
edit(236, "High-confidence unannotated splicing changes in every comparison, including the independently produced mapping-quality-filtered SH-SY5Y junction set.",
     "Stringent-filter unannotated splicing candidates in every comparison, including the mapping-quality-"
     "filtered SH-SY5Y junction set produced with a different extractor.", "S5.")
inv16b = ("**S16b.** Donor-level re-analysis of multiple sclerosis: sample- and donor-level Cliff’s δ for TRPC1 and the "
          "myelin markers in normal-appearing white matter, and for the marker-adjusted comparison.")
T.insert_after(P[248], inv16b, like=P[248])
edit(249, "The supplementary material also contains Figures S1–S8 and their legends.",
     "The supplementary material also contains Figures S1–S9 and their legends.", "S17.")

# ------------------------------------------------------------------ tables (text of captions, notes and header cells)
edit(259, "among the twelve core entry components examined event by event (Methods 2.3)",
     "among the twelve core SOCE-pathway genes examined event by event (Methods 2.3)", "Table 3.")
edit(259, "belongs to the wider nineteen-gene SOCE panel of Methods 2.9", "belongs to the nineteen-gene SOCE-associated set of Methods 2.9")
norm(P[259])
p262 = expect(262, "Note. Calls are unannotated")
edit(262, "The null column gives the high-confidence calls", "The null column gives the stringent-filter calls")
edit(262, "In the Tier 1, SOCE-machinery and positive-control columns NA means that no gene of that set carried a high-confidence call,",
     "In the Tier 1, SOCE-associated-set and positive-control columns NA means that no gene of that set carried a stringent-filter call,")
edit(262, "the high-confidence definition requires a novel splice site", "the stringent-filter definition requires a novel splice site")
norm(P[262])


def set_cell(cell, old, new):
    runs = [r for p in cell.paragraphs for r in p.runs]
    assert cell.text == old, (cell.text, old)
    runs[0].text = new
    for r in runs[1:]:
        r.text = ""


t1 = doc.tables[0]
hit = [r for r in t1.rows if r.cells[0].text == "SOCE regulators"]
assert len(hit) == 1
set_cell(hit[0].cells[0], "SOCE regulators", "Ca²⁺-entry regulators")
t4 = doc.tables[3]
for j, (old, new) in {3: ("High-confidence calls (genes)", "Stringent-filter calls (genes)"),
                      4: ("Positive controls, high-confidence", "Positive controls, stringent filter"),
                      5: ("Positive-control genes, high-confidence", "Positive-control genes, stringent filter"),
                      7: ("SOCE-machinery genes", "SOCE-associated-set genes")}.items():
    set_cell(t4.rows[0].cells[j], old, new)

# ------------------------------------------------------------------ figures: images and alt text
FIG = R / "figures" / "main"
T.replace_image_before(doc, "Figure 2.", FIG / "Figure2_calcium_responses.png", 6.05)
T.replace_image_before(doc, "Figure 4.", FIG / "Figure4_SOCE_splicing.png", 6.15)
T.replace_image_before(doc, "Figure 5.", FIG / "Figure5_ALS_expression.png", 6.05)
T.replace_image_before(doc, "Figure 6.", FIG / "Figure6_working_model.png", 6.05)
ALT = {
    "Figure 1.": "Bar charts with individual points: TARDBP mRNA in three groups, four target mRNAs in shTDP-43 versus "
                 "non-targeting control cells, and the WST-1 signal at 48 h.",
    "Figure 2.": "Panel A: two representative Fura-2 recordings on common axes, with a large readdition peak in the "
                 "control well and a small one in the shTDP-43 well. Panels B to D: ER release, readdition amplitude "
                 "and readdition-to-release ratio for three wells per group.",
    "Figure 3.": "Dot plots of adjusted transcript abundance and DESeq2 log2 fold changes for calcium-regulatory genes in "
                 "control and TDP-43-depleted SH-SY5Y libraries.",
    "Figure 4.": "Panel A: per-library PSI of four skipped-exon events with a schematic of PSI. Panel B: forest plot of "
                 "the 24-nucleotide STIM2 exon in six datasets. Panel C: CBARP exon 4 junction usage in SH-SY5Y and iPSC colonies.",
    "Figure 5.": "Heat map of Cliff's delta for TRPC1, SARAF and CBARP in ten ALS brain and spinal cord regions, with the "
                 "number of cases and controls in each row label.",
    "Figure 6.": "Diagram of laboratory measurements (RT-qPCR and Fura-2, from separate cultures) and three candidate "
                 "mechanisms suggested by public RNA-seq data, linked by dashed hypothesis arrows.",
}
for cap, text in ALT.items():
    idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.startswith(cap))
    pic = [p for p in doc.paragraphs[max(0, idx - 3):idx] if p._p.xpath(".//a:blip")]
    assert len(pic) == 1, cap
    T.set_alt_text(pic[0], text)

assert T.field_count(doc) == fields_before, (T.field_count(doc), fields_before)
MAIN_OUT = OUT / "manuscript/MANUSCRIPT_NEUROCHEMISTRY_INTERNATIONAL_FINAL.docx"
MAIN_OUT.parent.mkdir(parents=True, exist_ok=True)
# gene symbols that were left upright in passages about genes or transcripts (proteins stay upright)
_GENE_PARAS = ["Six comparisons with paired controls", "Cell-composition markers (", "SH-SY5Y cells were maintained in DMEM/F12",
               "Bioinformatic thresholds and correction families", "Mean TARDBP mRNA was", "Notably, the ALS increase in TRPC1",
               "S18. Cell-composition adjustment", "S18b. Cryptic STMN2", "S18c. Within the comparison group",
               "S18d. Donor-level NYGC sensitivity", "Note. SH-SY5Y, 0 versus 75 ng/mL doxycycline, three libraries per group"]
_n_ital = 0
for _op in _GENE_PARAS:
    _m = [q for q in doc.paragraphs if q.text.startswith(_op)]
    assert len(_m) == 1, (_op, len(_m))
    _before = _m[0].text
    _n_ital += T.italicise_in_place(_m[0])
    assert _m[0].text == _before
print("gene symbols made italic in place:", _n_ital)

doc.save(str(MAIN_OUT))
print("main manuscript written:", MAIN_OUT, "| citation fields:", fields_before, "->", T.field_count(doc))

# =========================================================================== supplementary material
sup = Document(str(SUPP_IN))
S = list(sup.paragraphs)


def sexpect(i, opening):
    assert S[i].text.startswith(opening), (i, S[i].text[:60], opening)
    return S[i]


sexpect(1, "TDP-43 knockdown is associated")
T.edit_text(S[1], OLD_TITLE, NEW_TITLE)
sexpect(7, "Supplementary results")
sexpect(8, "Supplementary figures")
for r in S[7]._p.findall(qn("w:r")):
    S[7]._p.remove(r)
for r in S[8]._p.findall(qn("w:r")):
    S[7]._p.append(copy.deepcopy(r))
runs7 = S[7]._p.findall(qn("w:r"))
runs7[0].find(qn("w:t")).text = "Supplementary results"
runs7[1].find(qn("w:t")).text = "1–6"
T.edit_text(S[8], "S1–S8", "S1–S9")
T.edit_text(sexpect(16, "S5."), "High-confidence unannotated splicing changes in every comparison, including the independently produced mapping-quality-filtered SH-SY5Y junction set.",
            "Stringent-filter unannotated splicing candidates in every comparison, including the mapping-quality-filtered SH-SY5Y junction set produced with a different extractor.")
T.insert_after(sexpect(28, "S16."), inv16b, like=S[28])

# Supplementary Results 3: coverage-based polyadenylation screen, details moved from Section 3.6
sexpect(47, "Supplementary Results 3.")
T.edit_text(S[47], "Supplementary Results 3. Polyadenylation screen in the two mouse lines",
            "Supplementary Results 3. Coverage-based polyadenylation screen: details")
apa1 = T.insert_after(S[47], (
    "**SH-SY5Y positive control.** Excluding CIGAR reference skips matters for the *STMN2* positive control: "
    "including them would have changed the *STMN2* intron 2 index difference from +0.145 to +0.489. With skips "
    "excluded, the *STMN2* intronic index shifted only moderately in the primary SH-SY5Y model, by +0.145 over all "
    "six libraries (bootstrap 95% CI +0.085 to +0.205). *STMN2* falls outside the depth-filtered summary table "
    "because one control library has a summed two-window depth of 2.97, just below the pre-specified threshold of "
    "3; one knockdown library has zero measured coverage in the 3′ window, which by the definition of the index "
    "contributes a value of 1.0 rather than missing data, and excluding that library gives +0.121 (+0.068 to "
    "+0.173) (Supplementary Figure S6B)."), like=S[48])
apa2 = T.insert_after(apa1, (
    "**SH-SY5Y units outside the SOCE-associated set.** Three terminal-exon units of the Ca²⁺ panel outside the "
    "nineteen-gene SOCE-associated set excluded zero: *ITPR3* (−0.152), *MICU3* (+0.097) and *ITPR1* (−0.055). "
    "Larger shifts occur in the cryptic positive controls analysed alongside the panel, the *UNC13A* intron 2 unit "
    "most of all (−0.267), which is how the assay is expected to behave in this model. Because the positive "
    "control responded only moderately, absence of a signal is not evidence that a gene is spared."), like=S[48])
T.insert_after(apa2, (
    "**iPSC-derived motor neurons.** Here 59 units in 29 genes passed the depth filter. The *STMN2* intron 2 index, "
    "which contains cryptic exon 2a, rose from 0.570 in controls to 0.819 in knockdown (Δ = +0.249, interval +0.208 "
    "to +0.290). The largest shifts among genes of the SOCE-associated set were *ATP2A2* intron 3 (−0.164) and, "
    "below 0.10, *SARAF* intron 5 (+0.097), *TRPC1* intron 1 (+0.083) and the *STIM2* terminal exon (−0.074). The "
    "SH-SY5Y *STIM1* intron 17 unit did not reach the depth filter here (*STIM2* intron 13 gave −0.024, interval "
    "spanning zero). With n = 2 per group the enumerated bootstrap has sixteen draws and its intervals are coarse, "
    "so the point estimates carry the information."), like=S[48])
T.edit_text(S[49], "No SOCE-machinery unit exceeded", "No unit of the SOCE-associated set exceeded")
T.rebuild(S[48], "**Mouse lines.** " + T.to_markup(S[48]))
T.rebuild(S[49], T.to_markup(S[49]))

# new Supplementary Results 5 and 6 after Supplementary Results 4
anchor = S[53]
h5 = T.insert_after(anchor, "Supplementary Results 5. Surrogate-variable sensitivity analysis (svaseq)", like=S[50])
h5.style = S[50].style
body5 = T.insert_after(h5, (
    "DESeq2 (v1.42.1) was refitted on the gene-level counts of the primary SH-SY5Y comparison (three control and "
    "three 75 ng/mL libraries; genes with at least 10 counts across the six libraries) with a design that included "
    "the two surrogate variables estimated by svaseq (sva v3.50.0) in addition to the group term "
    "(code/svaseq_sensitivity.R). Two surrogate variables reduced the residual degrees of freedom from four to two. "
    "Of the 14,012 genes with adjusted p values in both models, 1,694 met both differential-expression thresholds "
    "(q < 0.05 and |log2 fold change| ≥ 1) in the original model and 1,067 in the adjusted model; 850 met both in "
    "both models and all 850 kept their direction. The median absolute difference in log2 fold change was 0.20. "
    "Values for the genes named in the main text follow; the full comparison is in "
    "source_data/svaseq_sensitivity_SHSY5Y.csv."), like=S[51])
h6 = T.insert_after(body5, "Supplementary Results 6. STIM1 isoform-level testing", like=S[50])
h6.style = S[50].style
body6 = T.insert_after(h6, (
    "Two separate frameworks tested *STIM1* at isoform level. In the DRIMSeq–stageR analysis, the DRIMSeq feature "
    "table contained nine *STIM1* transcripts. Seven had DRIMSeq p values and entered the stageR confirmation "
    "stage, and all seven had a confirmation-stage adjusted p value of 1.0 (ENST00000300737.8, ENST00000526596.2, "
    "ENST00000532610.5, ENST00000698909.1, ENST00000698910.1, ENST00000698918.1 and ENST00000698919.1). The other "
    "two, ENST00000698912.1 and ENST00000698913.1, had missing DRIMSeq p values, were not eligible for stageR and "
    "were not evaluated (__source_data/STIM1_transcript_test_eligibility.csv__)."), like=S[51])
body6b = T.insert_after(body6, (
    "In the separate IsoformSwitchAnalyzeR/DEXSeq analysis, ENST00000698912.1 and ENST00000698913.1 were the two "
    "isoforms with significant usage changes (q = 1.6 × 10⁻¹⁰ and 5.5 × 10⁻³⁸), with small changes in isoform "
    "fraction (dIF +0.037 and −0.032); neither is predicted to carry a premature termination codon. The two "
    "predicted PTC isoforms, ENST00000698919.1 and ENST00000698918.1, did not change (q = 0.33 and 0.77). The two "
    "frameworks are different statistical procedures and are not a single confirmation chain. Because the two "
    "isoforms with significant usage changes could not be evaluated by stageR, the stage-wise procedure neither "
    "confirms nor refutes them. No confirmed PTC-associated isoform switch was identified in this model, which "
    "does not exclude one in other cells."), like=S[51])

# small table of the genes named in the main text (built after the paragraph, then moved into place)
import pandas as pd  # noqa: E402

sv = pd.read_csv(R / "source_data" / "svaseq_sensitivity_SHSY5Y.csv").set_index("gene")
order = ["STIM1", "TRPC1", "ORAI3", "SARAF", "CBARP", "ORAI1", "ATP2A3", "ATP2A2", "STIM2"]


def fmt_q(x):
    if x < 0.001:
        m, e = f"{x:.1e}".split("e")
        return f"{m} × 10{str(int(e)).translate(str.maketrans('0123456789-', '⁰¹²³⁴⁵⁶⁷⁸⁹⁻'))}"
    return f"{x:.3f}" if x < 0.1 else f"{x:.2f}"


tbl = sup.add_table(rows=1 + len(order), cols=5)
tbl.style = sup.styles["Table Grid"]
head = ["Gene", "log2FC, original model", "q, original model", "log2FC, with two surrogate variables", "q, with two surrogate variables"]
for c, h in zip(tbl.rows[0].cells, head):
    c.text = ""
    run = c.paragraphs[0].add_run(h)
    run.bold = True
    run.font.size = Pt(8.5)
for r, g in zip(tbl.rows[1:], order):
    v = sv.loc[g]
    vals = [g, f"{v.log2FC_published:+.2f}".replace("-", "−"), fmt_q(v.padj_published),
            f"{v.log2FC_with_SV:+.2f}".replace("-", "−"), fmt_q(v.padj_with_SV)]
    for cell, val in zip(r.cells, vals):
        cell.text = ""
        run = cell.paragraphs[0].add_run(val)
        run.italic = (val == g)
        run.font.size = Pt(8.5)
body5._p.addnext(tbl._tbl)

# figure legends
def scaption(i, opening, markup):
    p = sexpect(i, opening)
    T.rebuild(p, markup)
    return p


scaption(56, "Supplementary Figure S1.", (
    "**Supplementary Figure S1.** Example detection-power simulation for a three-versus-three design. Simulated "
    "power to detect a difference in PSI between two groups of three replicates at fixed depths of 10, 20, 50 and "
    "100 informative reads per sample and true ΔPSI values from 0.05 to 0.30. Replicate PSI was drawn from a normal "
    "distribution around 0.5 ± ΔPSI/2 (between-replicate standard deviation 0.05, limited to 0.01–0.99), reads were "
    "sampled binomially, and the groups were compared with a two-sample t test (equal variances) at nominal p < 0.05 "
    "without multiple-testing correction; 2,000 simulations per point. The simulation does not apply the rMATS "
    "model, FDR correction or |ΔPSI| threshold, so it illustrates the dependence of power on read depth and is not "
    "an estimate of rMATS power. The horizontal line marks 80% power and the vertical line marks |ΔPSI| = 0.10. "
    "Values are in source_data/power_simulation_S1.csv."))
T.edit_text(sexpect(64, "Supplementary Figure S3."), "and each track has its own y-axis.",
            "and each track has its own y-axis, so track heights do not allow a direct comparison of expression "
            "between groups or datasets.")
scaption(68, "Supplementary Figure S4.", (
    "**Supplementary Figure S4.** Positive controls and specificity of the annotation-free junction screen. (A) "
    "Literature cryptic-splicing positive controls recovered de novo in SH-SY5Y at 75 ng/mL doxycycline "
    "(stringent-filter calls, mapping-quality-filtered junction set), with knockdown/control junction-read counts. "
    "(B) Genes with a stringent-filter unannotated splicing change in each of four cumulative Ca²⁺ panels in the "
    "same call set; the numbers below the tiers are panel sizes, and the tiers are nested, not independent. (C) "
    "Number of literature positive-control genes recovered under the permissive definition in human comparisons "
    "(regtools junction set), including the FUS and TAF15 knockdown controls. Two controls after TDP-43 knockdown "
    "versus none after FUS or TAF15 knockdown in the same motor neurons is not a significant difference (Fisher’s "
    "exact test p = 0.48) and does not by itself establish specificity. Panels A–C use different thresholds and "
    "call sets, as stated in their headings. Mouse comparisons are omitted because the literature control events "
    "are human cryptic events."))
scaption(76, "Supplementary Figure S6.", (
    "**Supplementary Figure S6.** Coverage-based alternative-polyadenylation screen and *STMN2* control. (A) "
    "Depth-qualified SH-SY5Y intronic and distal 3′UTR index changes (units with |Δ| ≥ 0.05) with bootstrap 95% "
    "intervals; shaded rows are cryptic positive-control genes and unshaded rows are genes of the Ca²⁺ panel. Unit "
    "numbers are ordinal and are not canonical intron numbers, and the indices measure coverage gradients, not "
    "poly(A) sites. The *STIM2* intron 13 unit (†) contains the alternatively spliced *STIM2* exon and is not "
    "independent evidence for a polyadenylation event. (B) Coverage at both ends of *STMN2* intron 2 in each "
    "replicate after reference skips are excluded; the open marker is a measured zero at the 3′ end in one "
    "knockdown replicate."))
scaption(80, "Supplementary Figure S7.", (
    "**Supplementary Figure S7.** *TRPC1* expression across neurological comparison cohorts. (A) Cliff’s δ (case "
    "minus control) for ALS regions, other neurological disorders, Alzheimer’s disease, Parkinson’s disease and "
    "multiple sclerosis (two cohorts), grouped by disease; the number of cases and controls, or donors, is printed "
    "at the right of each row. Dark bars mark Benjamini–Hochberg q < 0.05 within each cohort’s comparison family "
    "and pale bars are not significant. (B) Re-analyses of the samples and donors of the GSE138614 "
    "multiple-sclerosis cohort (lesions, normal-appearing white matter and a comparison adjusted for *MBP*, *PLP1* "
    "and *GFAP*); they use overlapping samples and are not independent replicates, and daggers mark uncorrected "
    "p < 0.05 for the donor-level comparisons. Cohorts differ in tissue compartment, sample size and processing, so "
    "the display is descriptive across diseases and the values are not pooled."))
sexpect(84, "Supplementary Figure S8.")
T.edit_text(S[84], "(A) Percentage of ALS and non-neurological control samples with a cryptic STMN2 junction, by region, among samples with at least 20 junction reads at the donor; sample counts for each group are printed in the labels.",
            "(A) Percentage of ALS and non-neurological control samples with a detected cryptic STMN2 junction (at "
            "least one cryptic-junction read), by region, among samples with at least 20 junction reads at the donor; "
            "sample counts for each group are printed in the labels. Supplementary Table S18b uses a different "
            "threshold (cryptic PSI > 1%).")
T.rebuild(S[84], T.to_markup(S[84]))

# images
SF = R / "figures" / "supplementary"
for cap, png, width in (("Supplementary Figure S1.", "Supplementary_Figure_S1_detection_power.png", 5.7),
                        ("Supplementary Figure S4.", "Supplementary_Figure_S4_cryptic_controls.png", 5.2),
                        ("Supplementary Figure S6.", "Supplementary_Figure_S6_APA_coverage.png", 5.05),
                        ("Supplementary Figure S7.", "Supplementary_Figure_S7_cross_disease_TRPC1.png", 5.55),
                        ("Supplementary Figure S8.", "Supplementary_Figure_S8_STMN2_expression_vs_cryptic_PSI.png", 5.15)):
    T.replace_image_before(sup, cap, SF / png, width)

# Supplementary Figure S9 (new): copy the block of Supplementary Figure S8 (blank, heading, picture, caption)
s8_heading, s8_pic, s8_cap = S[82], S[83], S[84]
blank = copy.deepcopy(S[81]._p)
heading = copy.deepcopy(s8_heading._p)
pic = copy.deepcopy(s8_pic._p)
cap9 = copy.deepcopy(s8_cap._p)
s8_cap._p.addnext(blank)
blank.addnext(heading)
heading.addnext(pic)
pic.addnext(cap9)
h9, p9, c9 = Paragraph(heading, s8_cap._parent), Paragraph(pic, s8_cap._parent), Paragraph(cap9, s8_cap._parent)
T.edit_text(h9, "Supplementary Figure S8", "Supplementary Figure S9")
rid, _ = sup.part.get_or_add_image(str(SF / "Supplementary_Figure_S9_original_Fura2_recordings.png"))
blip = p9._p.xpath(".//a:blip")[0]
blip.set(qn("r:embed"), rid)
from PIL import Image  # noqa: E402

w, h = Image.open(SF / "Supplementary_Figure_S9_original_Fura2_recordings.png").size
cx = int(Inches(6.3))
cy = int(cx * h / w)
for e in p9._p.xpath(".//wp:extent") + p9._p.xpath(".//a:ext"):
    e.set("cx", str(cx))
    e.set("cy", str(cy))
for dp in p9._p.xpath(".//wp:docPr"):
    dp.set("id", "9")
    dp.set("name", "Picture 9")
    dp.set("title", "Supplementary Figure S9")
T.rebuild(c9, (
    "**Supplementary Figure S9.** Original Fura-2 recordings. The two recordings shown on common axes in Figure 2A, "
    "exported without redrawing from the GraphPad Prism trace project, each at the axis range of its own "
    "recording: the y axis spans 0–3 in A (non-targeting shRNA control) and 0.5–1.5 in B (shTDP-43), and the time "
    "windows differ, so the two panels are not directly comparable by eye. CPA (10 µM) and CaCl₂ (nominally 1.5 mM) "
    "additions are indicated. The recordings illustrate individual wells; group amplitudes were calculated from the "
    "original Prism measurements."))

# alt text of all supplementary figures (S5 still carried the earlier 'units of inference' wording)
SALT = {
    "Supplementary Figure S1.": "Line plot of simulated detection power against true delta PSI for four read depths, with the simulation assumptions printed below the plot.",
    "Supplementary Figure S2.": "Replicate-level PSI, bootstrap distribution and leave-one-out estimates for a TRPC1 skipped-exon call.",
    "Supplementary Figure S3.": "Sashimi-style coverage and junction reads at the CBARP exon 4 to exon 5 region in SH-SY5Y and iPSC colonies, and delta PSI of 32 CBARP rMATS events.",
    "Supplementary Figure S4.": "Bar charts of literature cryptic-splicing controls recovered in SH-SY5Y, stringent-filter events in cumulative calcium panels, and positive controls recovered in human comparisons.",
    "Supplementary Figure S5.": "Histogram of descriptive TDP-43 by NMD-inhibition interactions for genes with a cryptic junction and other genes, with CBARP marked; no inferential significance is assigned.",
    "Supplementary Figure S6.": "Forest plot of coverage-index changes in SH-SY5Y with shaded rows for cryptic positive-control genes, and STMN2 intron 2 coverage at both intron ends.",
    "Supplementary Figure S7.": "Bar charts of TRPC1 Cliff's delta by disease and region, and separate bars for multiple-sclerosis sensitivity analyses.",
    "Supplementary Figure S8.": "Bar chart of the percentage of samples with a detected cryptic STMN2 junction by region, and heat maps of Spearman correlations of selected genes with total STMN2 and cryptic STMN2 PSI.",
    "Supplementary Figure S9.": "Two original Fura-2 recordings, a control and a shTDP-43 well, each at its own axis range, with CPA and CaCl2 additions marked.",
}
for cap, text in SALT.items():
    idx = next(i for i, p in enumerate(sup.paragraphs) if p.text.startswith(cap))
    pics = [p for p in sup.paragraphs[max(0, idx - 3):idx] if p._p.xpath(".//a:blip")]
    assert len(pics) == 1, cap
    T.set_alt_text(pics[0], text)

SUPP_OUT = OUT / "supplementary/SUPPLEMENTARY_MATERIAL.docx"
SUPP_OUT.parent.mkdir(parents=True, exist_ok=True)
sup.save(str(SUPP_OUT))
print("supplement written:", SUPP_OUT)

# =========================================================================== highlights
hl = Document(str(HL_IN))
assert hl.paragraphs[3].text == "RNA analyses identified candidates without establishing a SOCE mechanism."
T.edit_text(hl.paragraphs[3], "RNA analyses identified candidates without establishing a SOCE mechanism.",
            "CBARP splicing changed in five of six TDP-43-depletion RNA-seq datasets.")
T.rebuild(hl.paragraphs[3], T.to_markup(hl.paragraphs[3]))
T.rebuild(hl.paragraphs[4], T.to_markup(hl.paragraphs[4]))  # gene symbols italic in the fourth item as well
for q in hl.paragraphs[1:]:
    assert len(q.text) <= 85, (len(q.text), q.text)
HL_OUT = OUT / "highlights_Neurochemistry_International.docx"
hl.save(str(HL_OUT))
print("highlights written:", HL_OUT, [len(q.text) for q in hl.paragraphs[1:]])
