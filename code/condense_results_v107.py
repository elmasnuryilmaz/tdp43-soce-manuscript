"""Shorten the body text without dropping results, numbers or supplementary pointers.

Sections 3.5 and 3.7 carried the repetition: the same caveat was stated two or three times,
the same qualification closed several consecutive paragraphs, and two sentences often carried
one fact. Every value, gene symbol, accession and supplementary pointer in the replaced text is
kept. Gene symbols are re-italicised afterwards, because replacing a whole paragraph collapses
its run formatting.
"""
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "manuscript/MANUSCRIPT_NEUROCHEMISTRY_INTERNATIONAL_FINAL.docx"

GENES = [
    "TARDBP", "STMN2", "UNC13A", "ACTL6B", "PFKP", "HDGFL2", "AGRN", "ARHGAP32", "ATG4B",
    "ELAVL3", "SETD5", "RSF1", "GPSM2", "KALRN", "CAMK2B", "POLDIP3", "SYNJ2", "TRPM3",
    "STIM1", "STIM2", "STIMATE", "SARAF", "CRACR2A", "CRACR2B", "ORAI1", "ORAI2", "ORAI3",
    "TRPC1", "ATP2A1", "ATP2A2", "ATP2A3", "CBARP", "SELENOK", "MCU", "MCUB", "MICU1",
    "MICU2", "MICU3", "MCUR1", "ITPR1", "ITPR2", "ITPR3", "RYR2", "RYR3", "SNAP25", "GFAP",
    "MBP", "PLP1", "MOG", "MAG", "XRN1", "UPF1", "SMG6", "ATP2B2", "ATP2B3", "CHGA",
    "Stmn2", "Saraf", "Atp2a2", "Trpc1", "Cbarp",
]

# paragraph index -> (unique opening of the current text, full replacement)
EDITS = {
    75: ("In control cells ORAI2 and ATP2A2 are the dominant members",
         "In control cells ORAI2 and ATP2A2 are the dominant members of their families (77% and 98% "
         "of the family total), whereas ORAI1 and ATP2A3 are minor members (15% and 1.5%); STIM1 and "
         "TRPC1 dominate their families (64% and 98%), and SARAF accounts for 82% of the "
         "SOCE-regulator pool. In knockdown the dominant ORAI2 and ATP2A2 transcripts fell modestly "
         "(DESeq2 log2FC −0.17 and −0.25; p_adj = 2.2 × 10⁻³ and 8.7 × 10⁻⁹), while ORAI1, ATP2A3, "
         "STIM1, TRPC1 and SARAF rose and ORAI3 rose about fivefold in adjusted TPM (DESeq2 log2FC "
         "2.06; p_adj = 2.3 × 10⁻⁶⁴). At family level the ORAI pool therefore grew (+27%) and changed "
         "in composition: ORAI3 went from 8% to 30% of ORAI transcripts and ORAI2 from 77% to 53%. "
         "The STIM (+48%) and SOCE-regulator (+30%) pools also rose, whereas the SERCA and "
         "mitochondrial-uptake pools fell (−12% and −25%), the latter mainly through MCU, MICU2 and "
         "MCUB."),
    94: ("rMATS and LeafCutter both constrain what can be found",
         "rMATS and LeafCutter both constrain what can be found, the first by the event classes and "
         "annotation it works from, the second by clustering rules that discard sparse junctions. To "
         "remove those constraints we extracted splice junctions directly from the alignments and "
         "tested local splicing variations with a beta-binomial model across eleven comparisons, "
         "repeating the primary SH-SY5Y comparison on a second, independently produced junction set "
         "(Methods 2.5)."),
    96: ("Specificity was tested by running the same analysis",
         "Specificity was tested by running the same analysis on FUS and TAF15 knockdown in the same "
         "iPSC-derived motor neurons, at the same depth and with the same design (Supplementary "
         "Figure S4C). Under the permissive definition the three comparisons produced call counts of "
         "the same order (141 for TDP-43, 124 for FUS and 126 for TAF15), but only TDP-43 knockdown "
         "recovered positive controls, two of the sixteen (STMN2 and KALRN). This dataset is too "
         "shallow for the high-confidence threshold to recover a control in any of the three "
         "comparisons (18, 26 and 23 events, none of them a control gene), and two controls against "
         "none is not by itself a significant difference (Fisher’s exact test p = 0.48), so the "
         "specificity of the procedure rests on the deeper comparisons, in which the same permissive "
         "analysis recovered thirteen of sixteen positive controls in SH-SY5Y and fifteen in iPSC "
         "colonies. The background call rate is therefore not specific to TDP-43, whereas the genes "
         "the analysis identifies are."),
    97: ("The null test also sets the limits of interpretation",
         "The null test also sets the limits of interpretation. Under the permissive rule, the "
         "null-to-real call ratios were 0.98 in iPSC colonies and 2.01 in K562 total RNA "
         "(Supplementary Table S14). The high-confidence null could be run in the three datasets with "
         "four control replicates: for every real call the split-control null produced 0.64 calls in "
         "iPSC colonies, 2.17 in K562 total RNA and 0.83 in mouse striatum (Supplementary Table "
         "S14). These ratios are not calibrated false-discovery rates, because the two analyses "
         "differ in sample size and in the number of tests, but they show that call counts are not "
         "interpretable on their own. In the remaining eight comparisons the controls could not be "
         "split, and positive-control recovery is the only available check. Across the eleven "
         "comparisons the burden of high-confidence unannotated splicing changes ranged from 12 to "
         "477 events, and of the three datasets with a null, the one with the largest burden also "
         "had the largest null."),
    98: ("Applied to the Ca²⁺ panels, the analysis returned",
         "Applied to the Ca²⁺ panels, the analysis returned an almost complete negative "
         "(Supplementary Table S7). The negative is informative where the analysis demonstrably "
         "worked: in SH-SY5Y at both doses and in K562 poly(A)+ mRNA, where positive controls were "
         "recovered at the high-confidence threshold, and less strongly in the TDP-43 knockdown of "
         "iPSC-derived motor neurons, where two were recovered under the permissive definition only. "
         "In K562 total RNA no positive control was recovered, and the three mouse comparisons have "
         "no conserved control, so the absence of calls there carries little information. Within the "
         "coverage, models and calling criteria used here, the only high-confidence unannotated "
         "change anywhere in the SOCE panel was the CBARP junction in iPSC colonies, called there "
         "together with TRPM3; STIM1, STIM2, ORAI1–3, TRPC1, SARAF, STIMATE and the SERCA and "
         "mitochondrial uptake genes carried none in any comparison. That single call comes from the "
         "dataset whose null test produced 0.64 calls for every real call, so it is not evidence on "
         "its own, although the CBARP junction is corroborated below. Permissive-tier and "
         "annotated-site events are detailed in Supplementary Table S7 and Figure S4B. This finding "
         "does not exclude lower-abundance or context-specific events, but it provides no support for "
         "a cryptic-splicing switch as the explanation for the observed SOCE phenotype."),
    101: ("FRASER, run on the nine-sample doxycycline series",
          "FRASER, run on the nine-sample doxycycline series, returned no genome-wide significant "
          "outlier and no difference in per-sample outlier burden between depleted and control "
          "libraries (4.8 versus 2.3 events at p < 10⁻⁵; p = 0.35). This is the expected behaviour of "
          "an outlier method in a cohort where the aberrant state is the majority, and we report it "
          "as a limit of that approach at this cohort size rather than as evidence against aberrant "
          "splicing."),
    109: ("Excluding CIGAR reference skips affected this positive control",
          "Excluding CIGAR reference skips matters for this positive control: including them would "
          "have changed the STMN2 intron 2 index difference from +0.145 to +0.489. With skips "
          "excluded, the STMN2 intronic index shifted only moderately in the primary SH-SY5Y model, "
          "by +0.145 over all six libraries (bootstrap 95% CI +0.085 to +0.205). STMN2 falls outside "
          "the depth-filtered summary table because one control library has a summed two-window "
          "depth of 2.97, just below the pre-specified threshold of 3; one knockdown library has zero "
          "measured coverage in the 3′ window, which by the definition of the index contributes a "
          "value of 1.0 rather than missing data, and excluding that library gives +0.121 (+0.068 to "
          "+0.173) (Supplementary Figure S6B)."),
    110: ("Among the depth-qualified units of the SOCE machinery",
          "Among the depth-qualified units of the SOCE machinery in the full-panel analysis, two "
          "intronic units showed moderate decreases (STIM1 intron 17, Δ = −0.173, bootstrap 95% CI "
          "−0.233 to −0.123; STIM2 intron 13, Δ = −0.156, −0.239 to −0.037). Neither is independent "
          "of the splicing results: the 5′ window of the STIM2 unit contains the alternatively "
          "spliced STIM2 exon of Table 3 (chr4:27,021,494–27,021,612), whose inclusion falls in "
          "knockdown and whose exon reads dominate that window (control index 0.96), so the decrease "
          "reports that splicing change, and the STIM1 unit is the intron immediately upstream of the "
          "alternatively included STIM1 exon of Table 3, ending 50 bp before it. Both are candidate "
          "coverage gradients rather than localised poly(A) sites. Sparse intronic windows produced "
          "unstable point estimates and were excluded by the pre-specified depth filter."),
    111: ("The analysis therefore identifies at most one moderate candidate gradient",
          "The analysis therefore identifies at most one moderate candidate gradient in the SOCE "
          "machinery, in STIM1, and does not establish a direct APA event in any of its genes. No "
          "index change of 0.05 or more was found for ORAI1–3, TRPC1, SARAF, STIMATE, CBARP or the "
          "SERCA genes, and the two units of those genes whose intervals exclude zero are negligible "
          "in size (SARAF terminal exon −0.036, ORAI2 terminal exon +0.012). Outside the SOCE "
          "machinery, three units of the Ca²⁺ panel likewise exclude zero: the ITPR3 terminal exon "
          "(−0.152), the MICU3 terminal exon (+0.097) and the ITPR1 terminal exon (−0.055), so the "
          "ER-release and mitochondrial-uptake arms show the same kind of candidate gradient as STIM1 "
          "rather than being spared. Larger shifts occur in the cryptic positive controls analysed "
          "alongside the panel, the UNC13A intron 2 unit most of all (−0.267), which is how the assay "
          "is expected to behave in this model. Because the positive control responded only "
          "moderately here, absence of a signal is not evidence that these genes are spared."),
    112: ("We therefore repeated the analysis in the iPSC-derived motor neurons",
          "We therefore repeated the analysis in the iPSC-derived motor neurons, where 59 units in 29 "
          "genes passed the same depth filter and the positive control behaved as intended: the index "
          "of STMN2 intron 2, which contains cryptic exon 2a, rose from 0.570 in controls to 0.819 in "
          "knockdown (Δ = +0.249, interval +0.208 to +0.290). Against that working positive control "
          "the SOCE-machinery genes again showed only small shifts, the largest being ATP2A2 intron 3 "
          "(−0.164) and, below 0.10, SARAF intron 5 (+0.097), TRPC1 intron 1 (+0.083) and the STIM2 "
          "terminal exon (−0.074). The SH-SY5Y STIM1 intron 17 unit did not reach the depth filter in "
          "this shallower dataset (STIM2 intron 13 gave −0.024, interval spanning zero). Here n = 2, "
          "so the enumerated bootstrap has sixteen draws and its intervals are coarse; the point "
          "estimates, not the intervals, carry the information."),
}


def replace_text(paragraph, new_text, tokens):
    """Rewrite a paragraph and re-italicise gene symbols, keeping the base run formatting."""
    import copy
    import re
    assert "fldChar" not in paragraph._p.xml and "fldSimple" not in paragraph._p.xml, "field-bearing paragraph"
    base = paragraph.runs[0]._r if paragraph.runs else None
    rpr = None if base is None else base.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr")
    pattern = re.compile(r"\b(" + "|".join(sorted(tokens, key=len, reverse=True)) + r")\b")
    pieces, pos = [], 0
    for m in pattern.finditer(new_text):
        if m.start() > pos:
            pieces.append((new_text[pos:m.start()], False))
        pieces.append((m.group(0), True))
        pos = m.end()
    if pos < len(new_text):
        pieces.append((new_text[pos:], False))
    for r in list(paragraph.runs):
        r._r.getparent().remove(r._r)
    for chunk, ital in pieces:
        run = paragraph.add_run()
        if rpr is not None:
            run._r.insert(0, copy.deepcopy(rpr))
        run.text = chunk
        run.italic = True if ital else None


# Paragraphs that carry Zotero fields are edited span by span, and every span is free of the
# gene symbols that are italic, so no citation field and no italic run is disturbed.
SPANS = {
    100: [("Measured as an rMATS event, inclusion of the exon was higher",
           "Measured as an rMATS event, inclusion was higher"),
          ("At junction level, inclusion of the SOAR exon (chr4:27,007,983–27,008,006) was positive "
           "in direction in all five TDP-43 comparisons in which it was measurable, but the effect "
           "was small and inconsistent. ΔPSI was +0.149 in SH-SY5Y",
           "At junction level the exon (chr4:27,007,983–27,008,006) was positive in direction in all "
           "five TDP-43 comparisons in which it was measurable, but small and inconsistent: ΔPSI "
           "+0.149 in SH-SY5Y"),
          ("because the two are not the same quantity. rMATS combines both flanking junctions and "
           "normalises by effective length. The junction-level PSI is the share of the single "
           "downstream junction among all junctions leaving that donor, which makes it the more "
           "sensitive and the noisier measure for a lowly used exon. On either measure the direction "
           "is reproducible and the magnitude is too small to support an isoform switch",
           "because rMATS combines both flanking junctions and normalises by effective length, "
           "whereas the junction-level PSI is the share of the single downstream junction among all "
           "junctions leaving that donor, the more sensitive and the noisier measure for a lowly used "
           "exon. On either measure the magnitude is too small to support an isoform switch")],
    113: [("in C2C12 and 131 in NSC34. Neither has a positive control for this assay: the ",
           "in C2C12 and 131 in NSC34, neither with a positive control for this assay, because the "),
          (" cryptic polyadenylation site is absent from the mouse gene",
           " cryptic polyadenylation site is absent from the mouse gene")],
    125: [("; in the second multiple sclerosis cohort the direction was the same but the study was "
           "underpowered (",
           "; the direction was the same in the second multiple sclerosis cohort, which was "
           "underpowered ("),
          ("although not after Benjamini–Hochberg correction (q = 0.10). Because several samples come "
           "from the same donor, we repeated the test with the donor as the unit of inference, which "
           "gave a larger difference (seven of ten",
           "although not after Benjamini–Hochberg correction (q = 0.10), and lower again with the "
           "donor as the unit of inference (seven of ten"),
          ("(donor-level δ = −1.000 for both). Multiple sclerosis showed marked astrogliosis (",
           "(donor-level δ = −1.000 for both), and multiple sclerosis showed marked astrogliosis (")],
}


def replace_span(paragraph, old, new):
    """Surgical substring replacement that leaves field codes and neighbouring runs intact."""
    nodes = paragraph._p.xpath(".//w:t")
    parts = [n.text or "" for n in nodes]
    whole = "".join(parts)
    if whole.count(old) != 1:
        raise ValueError(f"span not unique: {old[:60]!r} ({whole.count(old)})")
    start = whole.index(old)
    end = start + len(old)
    offsets, pos = [], 0
    for part in parts:
        offsets.append((pos, pos + len(part)))
        pos += len(part)
    first = next(i for i, (a, b) in enumerate(offsets) if a <= start < b)
    last = next(i for i, (a, b) in enumerate(offsets) if a < end <= b)
    a, _ = offsets[first]
    z, _ = offsets[last]
    if first == last:
        nodes[first].text = parts[first][: start - a] + new + parts[first][end - a:]
    else:
        nodes[first].text = parts[first][: start - a] + new
        for i in range(first + 1, last):
            nodes[i].text = ""
        nodes[last].text = parts[last][end - z:]


doc = Document(MANUSCRIPT)
P = doc.paragraphs
before = sum(len(p.text.split()) for p in P)
for idx, (opening, new_text) in EDITS.items():
    assert P[idx].text.startswith(opening), f"paragraph {idx} does not start with {opening!r}"
    old_n = len(P[idx].text.split())
    replace_text(P[idx], new_text, GENES)
    print(f"[{idx}] {old_n} -> {len(P[idx].text.split())} words")
for idx, spans in SPANS.items():
    old_n = len(P[idx].text.split())
    for old, new in spans:
        if old != new:
            replace_span(P[idx], old, new)
    print(f"[{idx}] {old_n} -> {len(P[idx].text.split())} words (field-bearing)")
after = sum(len(p.text.split()) for p in doc.paragraphs)
doc.save(MANUSCRIPT)
print(f"document total {before} -> {after} words ({before - after} removed)")
