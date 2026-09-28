"""Align supplementary and highlights wording with the single-plate Fura-2 design."""
from pathlib import Path
from docx import Document

root = Path(__file__).resolve().parents[1]
main = Document(root / "manuscript/MANUSCRIPT_NEUROCHEMISTRY_INTERNATIONAL_FINAL.docx")
title = main.paragraphs[0].text

supp_path = root / "supplementary/SUPPLEMENTARY_MATERIAL.docx"
supp = Document(supp_path)
old = supp.paragraphs[1].text
assert old == "TDP-43 knockdown is associated with reduced store-operated Ca²⁺ entry and altered calcium-regulatory RNA profiles in SH-SY5Y cells"
assert len(supp.paragraphs[1].runs) == 1
supp.paragraphs[1].runs[0].text = title
supp.save(supp_path)

highlight_path = root / "highlights_Neurochemistry_International.docx"
highlights = Document(highlight_path)
old = "No RNA-processing event in the SOCE machinery explained the functional change."
assert highlights.paragraphs[3].text == old
assert len(highlights.paragraphs[3].runs) == 1
highlights.paragraphs[3].runs[0].text = (
    "No RNA-processing event in the SOCE machinery explained the observed Fura-2 difference."
)
highlights.save(highlight_path)
print("Supplementary title and highlights aligned with manuscript")
