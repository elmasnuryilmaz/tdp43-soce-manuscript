"""Align the discussion, highlights and working-model figure with the one-plate Fura-2 design."""
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
src = (ROOT / "code/condense_results_v107.py").read_text()
exec(src[src.index("def replace_span("):src.index("doc = Document(MANUSCRIPT)")])

man = ROOT / "manuscript/MANUSCRIPT_NEUROCHEMISTRY_INTERNATIONAL_FINAL.docx"
doc = Document(man)
p = list(doc.paragraphs)

def change(i, old, new):
    assert old in p[i].text, (i, old[:70])
    replace_span(p[i], old, new)

change(5,
       "This single-plate Fura-2 observation suggests a lower calcium-readdition response after TDP-43 knockdown and altered calcium-regulatory RNA profiles, without evidence for a single causal RNA event.",
       "The single-plate Fura-2 observation requires independent replication; the RNA analyses identify candidate calcium-regulatory changes without establishing a causal link to that response.")
change(76, "provide candidate molecular explanations for the reduced SOCE phenotype.",
       "provide hypotheses for the lower calcium-readdition signal observed in one laboratory plate.")
change(131,
       "This study identifies a marked reduction in Fura-2-measured SOCE after TDP-43 knockdown and connects that functional phenotype to coordinated changes in calcium-regulatory transcripts. The central feature is a mismatch between abundance and function: the adjusted STIM and ORAI transcript pools increased, yet the SOCE amplitude decreased by approximately 84%.",
       "In one Fura-2 culture plate, TDP-43 knockdown wells had a lower calcium-readdition response. In an independent public SH-SY5Y model, adjusted STIM and ORAI transcript pools increased; these data generate hypotheses about the within-plate signal but do not establish a reproducible functional phenotype or its RNA mechanism.")
change(131,
       "This combination defines a multilayered calcium-homeostasis phenotype rather than uniform suppression of one channel or one RNA-processing event.",
       "Together, the observations motivate testing calcium regulation and RNA processing in matched, independently replicated experiments.")
change(143,
       "Blue boxes and solid arrows summarise what was measured in the laboratory cultures after TARDBP knockdown: higher TRPC1, STIM1, ORAI1 and ATP2A3 mRNA, a smaller Ca²⁺-readdition amplitude and a smaller, non-significant decrease in ER Ca²⁺ release.",
       "Blue boxes and solid arrows summarise laboratory measurements after TARDBP knockdown: higher TRPC1, STIM1, ORAI1 and ATP2A3 mRNA and, in three wells per group on one culture plate, lower mean Ca²⁺-readdition and ER-release amplitudes. The Fura-2 comparison is descriptive.")

# Replace only the embedded Figure 6 raster. Its figure-generation script has already been rerun.
png = ROOT / "figures/main/Figure6_working_model.png"
blip = p[142]._p.xpath(".//a:blip")[0]
rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
p[142].part.related_parts[rid]._blob = png.read_bytes()
w, h = Image.open(png).size
cx = int(Inches(6.05)); cy = int(cx * h / w)
for ext in p[142]._p.xpath(".//wp:extent") + p[142]._p.xpath(".//a:ext"):
    ext.set("cx", str(cx)); ext.set("cy", str(cy))
doc.save(man)

supp = ROOT / "supplementary/SUPPLEMENTARY_MATERIAL.docx"
d = Document(supp)
replace_span(d.paragraphs[12], "Fura-2 amplitudes", "Fura-2 amplitudes from three wells per group on one culture plate")
d.save(supp)

highlights = ROOT / "highlights_Neurochemistry_International.docx"
d = Document(highlights)
assert d.paragraphs[1].text.startswith("TDP-43 knockdown reduced")
replace_span(d.paragraphs[1], "TDP-43 knockdown reduced a SOCE-associated Ca²⁺ response in SH-SY5Y cells.",
             "One-plate Fura-2 measurements showed a lower Ca²⁺-readdition response after TDP-43 knockdown.")
d.save(highlights)

print('Updated manuscript, supplement and highlights')
