"""Remove remaining causal and culture-level wording after the Fura-2 unit correction."""
from pathlib import Path
from docx import Document

root = Path(__file__).resolve().parents[1]
src = (root / 'code/condense_results_v107.py').read_text()
exec(src[src.index('def replace_span('):src.index('doc = Document(MANUSCRIPT)')])
path = root / 'manuscript/MANUSCRIPT_NEUROCHEMISTRY_INTERNATIONAL_FINAL.docx'
doc = Document(path)
p = list(doc.paragraphs)

def change(i, old, new):
    assert old in p[i].text, (i, old)
    replace_span(p[i], old, new)

change(5, 'No confirmed RNA-processing event in the core SOCE machinery explained the functional change.',
       'No confirmed RNA-processing event in the core SOCE machinery explained the observed Fura-2 difference.')
change(67, 'The transcript changes and functional change are therefore in opposite directions.',
       'The measured mRNAs and the within-plate Fura-2 response therefore moved in opposite directions.')
change(132, 'Architecture of the functional phenotype.', 'Interpretation of the Fura-2 response.')
change(145, 'Normalising readdition to the release of the same culture',
       'Normalising readdition to the release of the same well')
doc.save(path)
print(path)
