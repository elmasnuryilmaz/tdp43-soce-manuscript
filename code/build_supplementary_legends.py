#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Regenerate supplementary/Supplementary_Figure_Legends.{md,docx} from the supplementary document (release v1.0.8).

The stand-alone legend files had been written for an earlier three-figure layout and had drifted from the
figures (an old four-gene Supplementary Figure S3, the four NMD conditions called 'units of inference').
They are now a copy of the captions in supplementary/SUPPLEMENTARY_MATERIAL.docx, which is the single
source: edit the captions there and run this script again.

Run from the repository root:  /usr/bin/python3 code/build_supplementary_legends.py
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx_edit_tools as T  # noqa: E402

R = Path(__file__).resolve().parents[1]
SRC = R / "supplementary" / "SUPPLEMENTARY_MATERIAL.docx"
MD = R / "supplementary" / "Supplementary_Figure_Legends.md"
DOCX = R / "supplementary" / "Supplementary_Figure_Legends.docx"

NOTE = ("These legends are copied from `supplementary/SUPPLEMENTARY_MATERIAL.docx` (release v1.0.8) by "
        "`code/build_supplementary_legends.py`; the figure files are in `figures/supplementary/`.")


def legends():
    doc = Document(str(SRC))
    out = []
    for p in doc.paragraphs:
        if re.match(r"Supplementary Figure S\d\.", p.text):
            mk = T.to_markup(p).replace("__", "")
            mk = re.sub(r"\*\*(Supplementary Figure S\d)\.\s*\*\*\s*", r"**\1.** ", mk)
            out.append(mk.strip())
    assert [re.match(r"\*\*Supplementary Figure (S\d)", m).group(1) for m in out] == [f"S{i}" for i in range(1, 10)], \
        "expected Supplementary Figures S1-S9 in order"
    return out


def add_runs(par, markup):
    """Bold (**) and italic (*) spans to runs; Arial 10.5 pt, as in the original file."""
    for tok in re.split(r"(\*\*.+?\*\*|\*.+?\*)", markup):
        if not tok:
            continue
        bold, ital = tok.startswith("**"), (tok.startswith("*") and not tok.startswith("**"))
        text = tok[2:-2] if bold else tok[1:-1] if ital else tok
        run = par.add_run(text)
        run.bold, run.italic = bold or None, ital or None
        run.font.name, run.font.size = "Arial", Pt(10.5)


def main():
    caps = legends()
    MD.write_text("# Supplementary figure legends\n\n" + NOTE + "\n\n" + "\n\n".join(caps) + "\n", encoding="utf-8")

    doc = Document(str(DOCX))
    for p in list(doc.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = doc.add_paragraph(style="First Paragraph")
    add_runs(first, NOTE.replace("`", ""))
    for cap in caps:
        add_runs(doc.add_paragraph(style="Body Text"), cap)
    doc.save(str(DOCX))
    print(f"written: {MD} and {DOCX} ({len(caps)} legends)")


if __name__ == "__main__":
    main()
