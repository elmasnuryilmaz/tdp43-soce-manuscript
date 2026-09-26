"""Second condensation pass: repeated conclusions in the Discussion and the two mouse/motor-neuron
APA paragraphs. Only sentences whose content is stated elsewhere are removed; no result, number or
citation is dropped. Spans inside field-bearing paragraphs contain no italic gene symbol.
"""
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "manuscript/MANUSCRIPT_NEUROCHEMISTRY_INTERNATIONAL_FINAL.docx"
src = (ROOT / "code/condense_results_v107.py").read_text()
exec(src[src.index("def replace_span("):src.index("doc = Document(MANUSCRIPT)")])

CUTS = {
    # Discussion: each removed sentence repeats a point made in full elsewhere
    133: [" Across independent datasets, canonical TDP-43 RNA targets behaved as expected, whereas "
          "calcium-gene processing changes were selective and context dependent."],
    136: [" These comparisons make the observed transcript pattern mechanistically informative "
          "without assigning the phenotype to one component: the likely variable is the assembled "
          "channel and its feedback kinetics in this cellular background."],
    137: ["The contrast is biologically meaningful because astroglial hyperactive entry, "
          "neuronal-like cell entry failure and ER-store depletion can all disturb calcium "
          "signalling through different routes. "],
    138: [" It nevertheless links the WST-1 phenotype to a broader bioenergetic and "
          "calcium-buffering response."],
    141: [" The resulting hierarchy is stronger than a long candidate list because it distinguishes "
          "recurrent signals from coverage-sensitive observations."],
    142: [" Together with the opposing direction of SOCE changes reported in SOD1 astrocytes, these "
          "findings indicate that cell identity and compensatory state shape the calcium-regulatory "
          "phenotype."],
}
SPANS = {
    112: [("passed the same depth filter and the positive control behaved as intended: the index of ",
           "passed the same depth filter and the positive control behaved as intended: the index of "),
          (" intron 2, which contains cryptic exon 2a, rose from 0.570 in controls to 0.819 in "
           "knockdown (Δ = +0.249, interval +0.208 to +0.290). Against that working positive control "
           "the ",
           " intron 2, which contains cryptic exon 2a, rose from 0.570 in controls to 0.819 in "
           "knockdown (Δ = +0.249, interval +0.208 to +0.290). Against it the "),
          (" intron 17 unit did not reach the depth filter in this shallower dataset (",
           " intron 17 unit did not reach the depth filter here ("),
          (" Here n = 2, so the enumerated bootstrap has sixteen draws and its intervals are coarse; "
           "the point estimates, not the intervals, carry the information.",
           " Here n = 2, so the enumerated bootstrap has sixteen draws and its intervals are coarse, "
           "and the point estimates carry the information.")],
    113: [(" No SOCE-machinery unit exceeded |Δ| = 0.30 in either line, and the two largest (",
           " No SOCE-machinery unit exceeded |Δ| = 0.30 in either line, and the two largest ("),
          (" We record it as a candidate rather than a finding: the NSC34 interval is wide, the "
           "myoblast line shows nothing there, the human and mouse units are matched by number "
           "rather than by sequence alignment, and a coverage gradient is not a poly(A) site.",
           " We record it as a candidate rather than a finding: the NSC34 interval is wide, the "
           "myoblast line shows nothing there, the units are matched by number rather than by "
           "sequence alignment, and a coverage gradient is not a poly(A) site.")],
}

doc = Document(MANUSCRIPT)
P = doc.paragraphs
before = sum(len(p.text.split()) for p in P)
for idx, sentences in CUTS.items():
    n = len(P[idx].text.split())
    for s in sentences:
        replace_span(P[idx], s, "")
    print(f"[{idx}] {n} -> {len(P[idx].text.split())} words")
for idx, spans in SPANS.items():
    n = len(P[idx].text.split())
    for old, new in spans:
        if old != new:
            replace_span(P[idx], old, new)
    print(f"[{idx}] {n} -> {len(P[idx].text.split())} words")
after = sum(len(p.text.split()) for p in doc.paragraphs)
doc.save(MANUSCRIPT)
print(f"document total {before} -> {after} words ({before - after} removed)")
