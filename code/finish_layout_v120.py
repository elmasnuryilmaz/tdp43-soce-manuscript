"""Final targeted wording and caption/inventory formatting after visual review."""
from pathlib import Path
from copy import deepcopy
import re
from docx import Document
from docx.text.run import Run
R=Path(__file__).resolve().parents[1]
h=(R/'code/condense_results_v107.py').read_text();exec(h[h.index('def replace_span('):h.index('doc = Document(MANUSCRIPT)')])
for path in ['manuscript/MANUSCRIPT_NEUROCHEMISTRY_INTERNATIONAL_FINAL.docx','supplementary/SUPPLEMENTARY_MATERIAL.docx']:
 d=Document(R/path)
 for p in d.paragraphs:
  a='Cervical and lumbar cord are well powered in this cohort, so the absence there is informative; thoracic cord and occipital cortex are not, since no gene reached significance in either.'
  if a in p.text:replace_span(p,a,'The lack of a statistically significant TRPC1 difference does not establish equivalence, particularly in the smaller regional comparisons.')
  m=re.match(r'(S\d+[a-d]?\.)(\s)',p.text)
  if m:
   for run in p.runs:run.bold=False
   run=next(run for run in p.runs if run.text)
   if run.text.startswith(m[1]):
    rest=run.text[len(m[1]):];run.text=m[1];run.bold=True
    if rest:
     q=Run(deepcopy(run._r),p);q.text=rest;q.bold=False;run._r.addnext(q._r)
  if p.text.startswith('Figure 6.'):
   for run in p.runs:run.italic=True
  for run in list(p.runs):
   if run.text.startswith('TRPC1 showed no statistically significant'):
    rest=run.text[5:];run.text='TRPC1';run.italic=True
    q=Run(deepcopy(run._r),p);q.text=rest;q.italic=False;run._r.addnext(q._r)
 d.save(R/path)
