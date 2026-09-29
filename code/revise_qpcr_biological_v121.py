"""One-time manuscript correction following the author's qPCR clarification."""
from pathlib import Path
from docx import Document
R=Path(__file__).resolve().parents[1]
h=(R/'code/condense_results_v107.py').read_text()
exec(h[h.index('def replace_span('):h.index('doc = Document(MANUSCRIPT)')])
changes={
'four technical RT-qPCR repeats per group; descriptive':'four biological RT-qPCR replicates per group',
'The laboratory observations require independent biological replication;':'The Fura-2 and WST-1 observations require independent biological replication;',
'four technical measurements of the same biological sample per group':'four biological replicates per group',
'These dates do not identify independent cultures or RNA isolations. ':'',
'All laboratory comparisons are descriptive. RT-qPCR comprises four biological replicates per group; the TARDBP and target-gene assays used separate RNA sets.':'RT-qPCR comprises four biological replicates per group; the TARDBP and target-gene assays used separate RNA sets. Two-sided Welch t-tests were applied to ΔCt values, with Holm correction across the four target genes and, separately, across the two TARDBP knockdown-versus-control comparisons. Adjusted p < 0.05 was the significance criterion. Relative-expression plots show means and SEM across biological replicates.',
'We report means, technical-repeat or well-to-well SEM, and observed differences without inferential p values, biological confidence intervals or standardised effect sizes. These measurements do not estimate between-experiment variability.':'For Fura-2 and WST-1, we report means, well-to-well SEM and observed differences without inferential p values, biological confidence intervals or standardised effect sizes. These two assays do not estimate between-experiment variability.',
'Mean TARDBP mRNA was 94.4% lower in the shTDP-43 sample than in the non-targeting shRNA sample and 94.8% lower than in the non-transduced sample (four technical RT-qPCR measurements per group; descriptive; Figure 1A).':'Mean TARDBP mRNA was 94.4% lower in shTDP-43 cells than in non-targeting shRNA controls and 94.8% lower than in non-transduced controls (four biological replicates per group; Holm-adjusted p = 8.3 × 10⁻¹¹ and 3.2 × 10⁻¹⁰, respectively; Figure 1A).',
'Relative to the non-targeting shRNA sample, mean measurements':'Relative to the non-targeting shRNA controls, mean measurements',
'four technical measurements per group; Figure 1B). These descriptive differences require biological replication.':'four biological replicates per group; Figure 1B). All four increases were significant in ΔCt tests after Holm correction (adjusted p = 0.0040, 0.0040, 0.0026 and 6.8 × 10⁻⁵, respectively).',
'Figure 1. Descriptive laboratory measurements':'Figure 1. Laboratory measurements',
'Each RT-qPCR group comprises four technical measurements of the same biological sample;':'Each RT-qPCR group comprises four biological replicates;',
'Bars show means with technical-repeat or well-to-well SEM; dots show individual measurements. No inferential significance tests are assigned to these comparisons.':'Bars show means with SEM across biological replicates in A and B and across wells in C; dots show individual measurements. RT-qPCR comparisons use two-sided Welch tests on ΔCt with Holm correction (Methods 2.16; Supplementary Table S1); WST-1 is descriptive.',
'RT-qPCR comprises four technical measurements per group; both the qPCR and Fura-2 comparisons are descriptive.':'RT-qPCR comprises four biological replicates per group; the Fura-2 comparison is descriptive.',
'RT-qPCR used four biological replicates per group and therefore does not provide independent biological confirmation of the RNA-seq changes. It was normalised':'RT-qPCR used four biological replicates per group and supported the direction of the four target-gene changes in an independent RNA sample set. It was normalised',
'including technical-repeat values and the descriptive summaries behind every laboratory figure panel':'including biological-replicate RT-qPCR values, well-level Fura-2 and WST-1 values, and the summaries behind every laboratory figure panel',
'four technical RT-qPCR measurements per group':'four biological RT-qPCR replicates per group',
'All laboratory summaries are descriptive.':'RT-qPCR summaries include biological-replicate SEM and Holm-adjusted ΔCt tests; Fura-2 and WST-1 summaries are descriptive.',
'release v1.0.6':'release v1.0.7',
'/releases/tag/v1.0.6':'/releases/tag/v1.0.7'
}
for file in ['manuscript/MANUSCRIPT_NEUROCHEMISTRY_INTERNATIONAL_FINAL.docx','supplementary/SUPPLEMENTARY_MATERIAL.docx']:
 d=Document(R/file)
 for p in d.paragraphs:
  for a,b in changes.items():
   if a in p.text: replace_span(p,a,b)
 d.save(R/file)
p=R/'highlights_Neurochemistry_International.docx'; d=Document(p)
for para in d.paragraphs:
 if para.text=='Technical RT-qPCR measurements showed higher calcium-regulatory mRNAs.':
  replace_span(para,para.text,'RT-qPCR in four biological replicates showed higher calcium-regulatory mRNAs.')
d.save(p)
print('qPCR biological-replicate manuscript correction applied')
