"""Add the Supplementary Results section that carries the analyses moved out of the main text.

Every number that left the manuscript in move_to_supplementary_v112.py is written here in full,
so nothing is lost: the FRASER run, the NMD interaction screen with its model and its tests, the
polyadenylation screen in the two mouse lines, and the multiple sclerosis analysis in detail.
"""
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SUPPLEMENTARY = ROOT / "supplementary/SUPPLEMENTARY_MATERIAL.docx"

NOTES = [
    ("Supplementary Results 1. Aberrant splicing outlier detection (FRASER)", [
        "FRASER v1.14.1 was run on the nine-sample SH-SY5Y doxycycline series (0, 25 and 75 ng/mL), "
        "counting split and non-split reads from the alignments and modelling ψ5, ψ3 and splicing "
        "efficiency θ (Methods 2.6).",
        "The run returned no genome-wide significant outlier and no difference in per-sample "
        "outlier burden between depleted and control libraries: 4.8 versus 2.3 events at "
        "p < 10⁻⁵, p = 0.35. This is the expected behaviour of an outlier method applied to a "
        "cohort in which the aberrant state is the majority, because the method treats aberrant "
        "splicing as a rare, sample-specific deviation from a cohort norm and two thirds of these "
        "samples are depleted. We therefore report the result as a limit of that approach at this "
        "cohort size rather than as evidence against aberrant splicing.",
    ]),
    ("Supplementary Results 2. Nonsense-mediated decay interaction", [
        "Gene-level counts were taken from GSE307054 (i3Neurons; TDP-43 knockdown crossed with "
        "knockdown of XRN1, UPF1 and SMG6 in four combinations, two replicates each) and "
        "normalised with the authors’ size factors. Because sequencing batch is confounded with "
        "TDP-43 status in that design, the interaction was formed as a difference of within-batch "
        "differences:",
        "interaction(c) = [log2(TDP-43 knockdown + NMD inhibition c) − log2(TDP-43 knockdown)] − "
        "[log2(control + NMD inhibition c) − log2(control)]",
        "for each of the four NMD-inhibition conditions c, giving four condition-level interaction "
        "values per gene; the two replicates within each condition are averaged. A positive "
        "interaction indicates a transcript that is produced upon TDP-43 loss and degraded by NMD. "
        "Genes with at least 10 normalised counts were tested (n = 19,145) with a four-condition "
        "one-sample t-test and an exact sign test.",
        "Because the four interventions reuse the same control and knockdown libraries, the eight "
        "difference-of-differences values that the design yields are not eight independent "
        "observations. We therefore took the four NMD-inhibition conditions as the unit of "
        "inference, averaging the two replicates within each condition and retaining the shared "
        "controls only within each contrast. Several positive-control and calcium-gene "
        "interactions were positive, but no gene passed genome-wide FDR correction in this "
        "conservative analysis (smallest q = 0.053).",
        "CBARP had the largest SOCE-panel interaction (+1.52 log₂; positive in all four "
        "conditions; Supplementary Table S10), but its conservative condition-level q-value was "
        "0.145 (two-sided t-test; exact one-sided sign-test q = 0.310). RYR2 and MCU also had "
        "positive interactions, without genome-wide significance.",
        "Two negative results follow. First, the prediction from isoform-level testing that STIM1 "
        "produces a premature-termination-codon-bearing, NMD-sensitive isoform was not supported: "
        "the interaction was −0.325 (p = 0.317, q = 0.508), in the opposite direction. Second, "
        "none of the four Ca²⁺ panels showed collective NMD sensitivity (one-sided Mann–Whitney "
        "p = 0.44, 0.91, 0.96 and 0.91 for Tiers 1–4), and the sixteen literature positive-control "
        "genes were not NMD-sensitive as a group either (median interaction −0.11 log₂; p = 0.62; "
        "Supplementary Table S10b). The last result matters for interpretation: genes that are "
        "known to produce cryptic, premature-termination-codon-bearing transcripts show no NMD "
        "sensitivity in this dataset, so the dataset does not demonstrate the sensitivity that "
        "would make a negative result for the SOCE machinery informative. The screen is therefore "
        "hypothesis-generating, and the main text uses it only to rank CBARP.",
    ]),
    ("Supplementary Results 3. Polyadenylation screen in the two mouse lines", [
        "The coverage-based polyadenylation screen of Methods 2.7 gave 74 qualifying units in "
        "C2C12 and 131 in NSC34. Neither line has a positive control for this assay, because the "
        "STMN2 cryptic polyadenylation site is absent from the mouse gene (Melamed et al., 2019), "
        "so the Stmn2 units cannot serve as one.",
        "No SOCE-machinery unit exceeded |Δ| = 0.30 in either line, and the two largest (Atp2a2 "
        "intron 6, +0.285, and Trpc1 intron 7, +0.260, both in C2C12) have intervals that include "
        "zero and are not reproduced in the other line. One unit is positive in both motor-neuron "
        "models: the SARAF intron 5 index rises in the iPSC-derived motor neurons (+0.097) and the "
        "Saraf intron 5 index in NSC34 (+0.204), while in SH-SY5Y the same unit is uninformative "
        "(0.000, interval −0.264 to +0.241). We record it as a candidate rather than a finding: "
        "the NSC34 interval is wide, the myoblast line shows nothing there, the human and mouse "
        "units are matched by number rather than by sequence alignment, and a coverage gradient is "
        "not a poly(A) site. All values are in Supplementary Table S11.",
    ]),
    ("Supplementary Results 4. Multiple sclerosis in detail", [
        "TRPC1 was decreased in multiple sclerosis at donor level across all sampled lesion types "
        "(δ = −0.840; q = 0.038) and in lesions at sample level (δ = −0.594; q = 1.3 × 10⁻⁴). In "
        "the second cohort (GSE123496) the direction was the same but the study was underpowered "
        "(five regions pooled, δ = −0.226; q = 0.49).",
        "Because TRPC1 contributes to SOCE in oligodendrocyte precursor cells (Paez et al., 2011), "
        "we asked whether the decrease simply reflected demyelination and loss of "
        "oligodendrocyte-lineage cells. In normal-appearing white matter, where the myelin markers "
        "MBP (δ = +0.051), PLP1 (−0.074), MOG (−0.257) and MAG (−0.299) were unchanged, TRPC1 was "
        "lower at sample level (δ = −0.482, p = 0.005), although not after Benjamini–Hochberg "
        "correction (q = 0.10), and lower again with the donor as the unit of inference (seven of "
        "ten multiple sclerosis donors with normal-appearing white matter samples versus five "
        "control donors, δ = −0.771, uncorrected p = 0.030).",
        "In the full GSE138614 comparison, regression of TRPC1 on MBP, PLP1 and GFAP attenuated "
        "the difference at both levels: from δ = −0.562 to Cliff’s δ of −0.418 on the residuals at "
        "sample level (p = 0.002), and from δ = −0.840 to −0.640 at donor level, where the "
        "adjusted difference did not reach significance (p = 0.055). TRPC1 was lower in every "
        "lesion type, including remyelinating and inactive lesions (donor-level δ = −1.000 for "
        "both), and multiple sclerosis showed marked astrogliosis (GFAP δ = +0.84 to +0.96) with "
        "TRPC1 still falling. All values are in Supplementary Table S16 and the donor-level "
        "re-analysis in Supplementary Table S16b.",
    ]),
]

doc = Document(SUPPLEMENTARY)
anchor = next(p for p in doc.paragraphs if p.text.strip() == "Supplementary Figure S1")

section = doc.add_paragraph("Supplementary results", style="Heading 1")
intro = doc.add_paragraph(
    "The analyses below were carried out as described in the Methods and are reported here in "
    "full. They are placed in the supplementary material because each is limited in a way that is "
    "stated in its own text: the outlier method cannot work in a cohort where the perturbed state "
    "is the majority, the nonsense-mediated decay dataset confounds batch with TDP-43 status and "
    "shows no sensitivity even in the literature positive controls, the mouse polyadenylation "
    "screen has no available positive control, and the multiple sclerosis analysis is a "
    "cross-disease comparison rather than part of the ALS argument.")
anchor._p.addprevious(section._p)
anchor._p.addprevious(intro._p)
for title, paragraphs in NOTES:
    h = doc.add_paragraph(title, style="Heading 1")
    anchor._p.addprevious(h._p)
    for text in paragraphs:
        body = doc.add_paragraph(text)
        anchor._p.addprevious(body._p)

contents = next(p for p in doc.paragraphs if p.text.startswith("Supplementary figures"))
extra = doc.add_paragraph("Supplementary results\t1–4")
contents._p.addprevious(extra._p)

doc.save(SUPPLEMENTARY)
print("supplementary results added")
