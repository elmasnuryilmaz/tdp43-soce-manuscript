"""Apply the 25 September 2026 audit corrections to the live DOCX files.

Run once, after revise_manuscript_v105.py. Text edits operate on Word text nodes, so the
embedded Zotero citation fields are untouched.

  1  Figure 1 is cited in Section 3.1 (it was cited nowhere in the running text)
  2  the cryptic-STMN2 proxy correlations are reported again in Section 3.9 (Supplementary
     Table S9 had become an uncited file after commit 34255c9)
  3  the myelin/glia-adjusted multiple sclerosis result is stated correctly for both levels
  4  Supplementary Table S11 now holds every depth-qualified core unit, so the small values
     quoted in Section 3.7 can be checked; its description says so
  5  the closing sentence of Section 3.5 no longer contradicts the CBARP exception it states
  6  the APA sentence says which gene set the three units belong to
  7  the enrichment conclusion follows the permutation p value rather than "largely attributable"
  8  the specificity claim rests on the deep datasets, not on 2/16 versus 0/16
  9  the supplementary document calls S12 the machine-readable version of Table 4, not Table 3
 10  "three independent cohorts" becomes three diseases in four datasets
 11  "both human models" names the two models compared at junction level
 12  the twelve genes examined event by event are "core entry components"; "SOCE machinery"
     keeps the nineteen-gene meaning defined in Methods 2.9
"""
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "manuscript/MANUSCRIPT_NEUROCHEMISTRY_INTERNATIONAL_FINAL.docx"
SUPPLEMENTARY = ROOT / "supplementary/SUPPLEMENTARY_MATERIAL.docx"


def replace(paragraph, old, new):
    nodes = paragraph._p.xpath(".//w:t")
    parts = [node.text or "" for node in nodes]
    whole = "".join(parts)
    if whole.count(old) != 1:
        raise ValueError(f"Expected exactly one occurrence of {old!r}; got {whole.count(old)}")
    start = whole.index(old)
    end = start + len(old)
    offsets, position = [], 0
    for part in parts:
        offsets.append((position, position + len(part)))
        position += len(part)
    first = next(i for i, (a, b) in enumerate(offsets) if a <= start < b)
    last = next(i for i, (a, b) in enumerate(offsets) if a < end <= b)
    a, _ = offsets[first]
    z, _ = offsets[last]
    if first == last:
        nodes[first].text = parts[first][: start - a] + new + parts[first][end - a:]
    else:
        nodes[first].text = parts[first][: start - a] + new
        for i in range(first + 1, last):
            nodes[i].text = ""
        nodes[last].text = parts[last][end - z:]


doc = Document(MANUSCRIPT)
p = list(doc.paragraphs)
assert p[64].text.startswith("Lentiviral shRNA reduced") and p[130].text.startswith("Gene-level")

# 1 — Figure 1
replace(p[64], "p < 0.0001, n = 4).", "p < 0.0001, n = 4; Figure 1A).")
replace(p[65], "(adjusted p = 0.00038; n = 4, two-tailed t-tests)",
        "(adjusted p = 0.00038; n = 4, two-tailed t-tests; Figure 1B)")
replace(p[68], "a 38.5% decrease (n = 4 wells; descriptive).",
        "a 38.5% decrease (n = 4 wells; descriptive; Figure 1C).")

# 12 — the twelve core entry components
replace(p[22], "the coverage criteria in twelve SOCE-machinery genes",
        "the coverage criteria in the twelve core entry components")
replace(p[41], "CBARP was included in this fixed panel for the subsequent cross-dataset event ranking.",
        "CBARP was included in this fixed panel for the subsequent cross-dataset event ranking. "
        "The twelve genes examined event by event in Section 2.3 are a subset of these nineteen "
        "and are called the core entry components.")
replace(p[87], "to the twelve SOCE-machinery genes examined event by event",
        "to the twelve core entry components examined event by event")
replace(p[92], "Per-library PSI for the four SOCE-machinery skipped-exon events of Table 3",
        "Per-library PSI for the four skipped-exon events of Table 3")
replace(p[255], "among the twelve SOCE-machinery genes examined event by event (Methods 2.3), with "
                "bootstrap confidence intervals; all four are skipped-exon events. CBARP was not "
                "among the twelve; its events are in Supplementary Table S3.",
        "among the twelve core entry components examined event by event (Methods 2.3), with "
        "bootstrap confidence intervals; all four are skipped-exon events. CBARP belongs to the "
        "wider nineteen-gene SOCE panel of Methods 2.9 but not to these twelve; its events are in "
        "Supplementary Table S3.")

# 11 — which human models
replace(p[89], "In both human models the canonical junction nearly disappeared",
        "In the two human models compared at junction level the canonical junction nearly disappeared")

# 8 — specificity
replace(p[96], "In the two deeper TDP-43 comparisons the same permissive analysis recovered thirteen "
               "of sixteen positive controls in SH-SY5Y and fifteen in iPSC colonies. The background "
               "rate of the procedure is not specific to TDP-43; the genes it identifies are.",
        "Two controls against none is not by itself a significant difference (Fisher's exact test "
        "p = 0.48), so the specificity of the procedure rests on the deeper comparisons, in which "
        "the same permissive analysis recovered thirteen of sixteen positive controls in SH-SY5Y "
        "and fifteen in iPSC colonies. The background call rate is therefore not specific to "
        "TDP-43, whereas the genes the analysis identifies are.")

# 5 — closing sentence of 3.5
replace(p[98], "Within the coverage, models and calling criteria used here, we did not detect "
               "high-confidence unannotated splicing changes in the store-operated entry machinery.",
        "Within the coverage, models and calling criteria used here, the only high-confidence "
        "unannotated change anywhere in the SOCE panel was the CBARP junction in iPSC colonies, and "
        "none of the core entry components carried one in any comparison.")

# 6 — which gene set the three units belong to
replace(p[111], "Outside the SOCE machinery three units likewise exclude zero: the ITPR3 terminal "
                "exon (−0.152), the MICU3 terminal exon (+0.097) and the ITPR1 terminal exon "
                "(−0.055), so the ER-release and mitochondrial-uptake arms show the same kind of "
                "candidate gradient as STIM1 rather than being spared.",
        "Outside the SOCE machinery, three units of the Ca²⁺ panel likewise exclude zero: the ITPR3 "
        "terminal exon (−0.152), the MICU3 terminal exon (+0.097) and the ITPR1 terminal exon "
        "(−0.055), so the ER-release and mitochondrial-uptake arms show the same kind of candidate "
        "gradient as STIM1 rather than being spared. Larger shifts occur in the cryptic positive "
        "controls analysed alongside the panel, the UNC13A intron 2 unit most of all (−0.267), "
        "which is the expected behaviour of the assay in this model.")

# 7 — enrichment conclusion
replace(p[117], "The apparent enrichment in this model is therefore largely attributable to the "
                "length and expression properties of Ca²⁺ genes.",
        "Against that null the excess is not significant, so this model provides no evidence of "
        "enrichment beyond the length and expression properties of Ca²⁺ genes.")

# 3 and 10 — multiple sclerosis paragraph
replace(p[125], "Three independent cohorts extended this (Supplementary Figure S7).",
        "Three independent diseases, in four datasets, extended this (Supplementary Figure S7).")
replace(p[125], "In the full GSE138614 sample-level comparison, regression of TRPC1 on MBP, PLP1 and "
                "GFAP attenuated the unadjusted difference from δ = −0.562 to Cliff’s δ of −0.418 on "
                "the residuals (p = 0.002) but not at donor level (δ = −0.640, p = 0.055).",
        "In the full GSE138614 comparison, regression of TRPC1 on MBP, PLP1 and GFAP attenuated the "
        "difference at both levels: from δ = −0.562 to Cliff’s δ of −0.418 on the residuals at "
        "sample level (p = 0.002), and from δ = −0.840 to −0.640 at donor level, where the adjusted "
        "difference was no longer significant (p = 0.055).")

# 2 — the proxy correlations return to the Results
proxy = doc.add_paragraph(style="Body Text")
proxy.add_run(
    "Within ALS samples the proxy did not track the three candidate transcripts. Across the ten "
    "regions, correlations of TRPC1, SARAF and CBARP with cryptic STMN2 PSI ranged from ρ = −0.29 "
    "to +0.18, and only one of these thirty tests reached significance (SARAF in lumbar cord, "
    "ρ = −0.194, q = 0.049), in the direction opposite to the increase seen in ALS tissue. This "
    "exploratory analysis therefore does not support attributing the regional expression "
    "differences to TDP-43-dependent RNA processing; all tested correlations are provided in "
    "Supplementary Table S9.")
p[130]._p.addnext(proxy._p)

# 4 — what S11 now contains
replace(p[243], "S11. Depth-qualified intronic polyadenylation and 3′UTR usage estimates in all four "
                "comparisons, with the genomic windows of every unit; these are coverage gradients, "
                "not direct poly(A)-site calls.",
        "S11. Depth-qualified intronic polyadenylation and 3′UTR usage estimates in all four "
        "comparisons, with the genomic windows of every unit; these are coverage gradients, not "
        "direct poly(A)-site calls. The candidate_gradient column marks the units with |Δ| ≥ 0.05 "
        "whose interval excludes zero; the remaining rows are the other depth-qualified units of "
        "the core genes and of the cryptic positive controls.")

doc.save(MANUSCRIPT)

# 9 and 4 — supplementary document
sup = Document(SUPPLEMENTARY)
sp = sup.paragraphs
s12 = next(x for x in sp if x.text.startswith("S12. Machine-readable version of Table 3"))
replace(s12, "Machine-readable version of Table 3", "Machine-readable version of Table 4")
s11 = next(x for x in sp if x.text.startswith("S11. Depth-qualified"))
replace(s11, "these are coverage gradients, not direct poly(A)-site calls.",
        "these are coverage gradients, not direct poly(A)-site calls. The candidate_gradient column "
        "marks the units with |Δ| ≥ 0.05 whose interval excludes zero; the remaining rows are the "
        "other depth-qualified units of the core genes and of the cryptic positive controls.")
sup.save(SUPPLEMENTARY)
print("revised:", MANUSCRIPT.name, "and", SUPPLEMENTARY.name)


def italicise(paragraph, tokens):
    """Rebuild a paragraph's runs so that each gene symbol is italic.

    Only for paragraphs without field codes; the base run formatting is copied from the first run.
    """
    import copy
    import re as _re
    assert "fldChar" not in paragraph._p.xml and "fldSimple" not in paragraph._p.xml
    text = paragraph.text
    base = paragraph.runs[0]._r if paragraph.runs else None
    pattern = _re.compile(r"\b(" + "|".join(_re.escape(t) for t in sorted(tokens, key=len, reverse=True)) + r")\b")
    pieces, pos = [], 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            pieces.append((text[pos:m.start()], False))
        pieces.append((m.group(0), True))
        pos = m.end()
    if pos < len(text):
        pieces.append((text[pos:], False))
    for r in list(paragraph.runs):
        r._r.getparent().remove(r._r)
    for chunk, ital in pieces:
        run = paragraph.add_run()
        if base is not None:
            rpr = base.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr")
            if rpr is not None:
                run._r.insert(0, copy.deepcopy(rpr))
        run.text = chunk
        run.italic = True if ital else None
    return len(pieces)


if __name__ == "__main__" or True:
    # gene symbols lost their italics where a whole sentence was replaced above
    doc2 = Document(MANUSCRIPT)
    q = doc2.paragraphs
    italicise(q[111], ["ITPR1", "ITPR3", "MICU3", "STIM1", "STIM2", "SARAF", "ORAI2", "ORAI1", "ORAI3",
                       "TRPC1", "STIMATE", "CBARP", "UNC13A"])
    italicise(q[131], ["TRPC1", "SARAF", "CBARP", "STMN2"])
    italicise(q[256], ["CBARP"])
    doc2.save(MANUSCRIPT)
    print("italics restored in paragraphs 111, 131 and 256")
