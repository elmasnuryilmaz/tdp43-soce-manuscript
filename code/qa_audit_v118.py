"""Independent regression checks for the September 29 scientific corrections."""
from pathlib import Path
import io,subprocess,zipfile
import numpy as np
import pandas as pd
import openpyxl
from lxml import etree
from docx import Document
R=Path(__file__).resolve().parents[1]
passed=[]
def check(label,ok):
    if not ok: raise AssertionError(label)
    passed.append(label)
s=pd.read_csv(R/'supplementary/S13_STIM2_SOAR_exon_junction_level.csv')
check('S13 includes 16 available junctions in nine comparisons',len(s)==16 and s.flank.value_counts().to_dict()=={'upstream':8,'downstream':8} and s.comparison.nunique()==9)
check('S13 includes both mouse cell models',{'C2C12','NSC34'}<=set(s.comparison))
check('S13 downstream denominators are acceptor LSVs',s[s.flank=='downstream'].lsv_type.eq('acceptor').all())
for mouse in [False,True]:
 z=s[s.comparison.isin(['C2C12','NSC34'])==mouse]
 check(f'S13 exon boundaries match source convention (mouse={mouse})',z[z.flank=='upstream'].end.eq(54110115 if mouse else 27007982).all() and z[z.flank=='downstream'].start.eq(54110140 if mouse else 27008007).all())
check('C2C12 effect is negative on both flanks',s[s.comparison=='C2C12'].delta_PSI.lt(0).all())
st=pd.read_csv(R/'source_data/STIM1_transcript_test_eligibility.csv')
check('STIM1 nine features, seven evaluated',len(st)==9 and st.stageR_eligible.sum()==7)
check('STIM1 untested features are the separately reported DEXSeq features',set(st.loc[~st.stageR_eligible,'feature_id'])=={'ENST00000698912.1','ENST00000698913.1'})
check('STIM1 seven confirmations equal one and missing features remain missing',st.loc[st.stageR_eligible,'stageR_confirmation_adjusted_p'].eq(1).all() and st.loc[~st.stageR_eligible,'stageR_confirmation_adjusted_p'].isna().all())
n=pd.read_csv(R/'source_data/nmd_descriptive_all_genes.csv.gz')
pd10=pd.read_csv(R/'supplementary/S10_NMD_interaction_SOCE_panel.csv')
pd10b=pd.read_csv(R/'supplementary/S10b_NMD_panel_descriptive_summary.csv')
for name,df in [('full NMD',n),('S10',pd10),('S10b',pd10b)]:
 check(name+' excludes inferential p/q columns',not any(c.lower().startswith(('p_','q_','padj','pval')) or c.lower() in {'p','q'} for c in df))
check('NMD expression filter retains 19145 genes',len(n)==19145)
cols=['interaction_XRN1','interaction_XRN1_SMG6','interaction_XRN1_UPF1','interaction_UPF1_SMG6']
check('NMD reported means are the four-contrast averages',np.allclose(n[cols].mean(axis=1),n.interaction_log2))
check('NMD reported ranges retain the condition variability',np.allclose(n[cols].min(axis=1),n.min_condition_interaction) and np.allclose(n[cols].max(axis=1),n.max_condition_interaction))
pd.testing.assert_frame_equal(pd10.set_index('ensembl_gene_id'),n.set_index('ensembl_gene_id').loc[pd10.ensembl_gene_id],check_exact=False)
check('S10 is an exact selection from full descriptive results',True)
# Directly recompute CBARP from count libraries, independent of the table builder.
base=R.parent/'07_DISK_ANALIZLERI/nmd_GSE307054'
raw=pd.read_csv(base/'GSE307054_counts.csv.gz',index_col=0)
meta=pd.read_csv(base/'GSE307054_meta.csv.gz',index_col=0)
x=np.log2(raw.loc['ENSG00000099625',meta.index]/meta.sizeFactor+1)
vals=[]
for c in ['X','XS','XU','US']:
 vals.append(x[[f'T{c}_1',f'T{c}_2']].mean()-x[[f'C{c}_1',f'C{c}_2']].mean()-x[['TDPKD_1','TDPKD_2']].mean()+x[['CON_1','CON_2']].mean())
check('CBARP contrasts reproduce directly from normalized libraries',np.allclose(vals,n[n.gene=='CBARP'][cols].iloc[0].values.astype(float)))
old=openpyxl.load_workbook(io.BytesIO(subprocess.check_output(['git','show','a80c10e:supplementary/S1_laboratory_source_data.xlsx'],cwd=R)))
new=openpyxl.load_workbook(R/'supplementary/S1_laboratory_source_data.xlsx')
allowed={'README':{'B2','B3','B4','B8'},'TARDBP_qPCR':{'B1'},'Target_qPCR_Ct':{'C1'},'Target_qPCR_rel':{'C1'},'Summary_stats':{f'{c}{i}' for c in 'GH' for i in range(2,14)}}
check('S1 preserves all source worksheets',old.sheetnames==new.sheetnames)
for name in old.sheetnames:
 for row in old[name]:
  for cell in row:
   if cell.coordinate not in allowed.get(name,set()):
    check(f'S1 preserved {name}!{cell.coordinate}',cell.value==new[name][cell.coordinate].value)
check('S1 removes qPCR inferential p values',all(new['Summary_stats'][f'H{i}'].value is None for i in range(2,14)))
donor=pd.read_csv(R/'supplementary/S18d_NYGC_donor_level_sensitivity.csv')
z=donor[(donor.group=='ALS')&(donor.region=='Cerebellum')&(donor.gene=='TRPC1')]
check('S18d distinguishes 158 cerebellar samples from 147 donors',len(z)==5 and z.n_case_samples.eq(158).all() and z.n_case_donors.eq(147).all())
check('Cerebellar donor sensitivity has positive effects in all models',z.cliffs_delta.gt(0).all() and z.q_value.lt(.05).all())
paths=['manuscript/MANUSCRIPT_NEUROCHEMISTRY_INTERNATIONAL_FINAL.docx','supplementary/SUPPLEMENTARY_MATERIAL.docx']
texts=[]
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
for path in paths:
 d=Document(R/path); text='\n'.join(p.text for p in d.paragraphs); texts.append(text)
 for suffix in ['S18.','S18b.','S18c.','S18d.']:
  check(path+' inventory '+suffix,suffix in text)
 check(path+' no obsolete release', 'v1.0.5' not in text)
 check(path+' labels RT-qPCR as technical','technical' in text and 'RT-qPCR' in text)
root=etree.fromstring(zipfile.ZipFile(R/paths[0]).read('word/document.xml'))
check('Main manuscript retains 79 Zotero instruction nodes',len([t for t in root.xpath('//w:instrText/text()',namespaces=ns) if 'ZOTERO' in t])==79)
check('Main no independent culture inference','three independent cultures' not in texts[0] and 'principal findings were unchanged' not in texts[0])
check('Main states technical-replicate limitation','biological replication' in texts[0])
for f in ['figures/main/Figure1_functional_consequences.svg','figures/graphical_abstract.svg']:
 svg=(R/f).read_text()
 check(f+' no qPCR inferential labels',not any(v in svg for v in ['p &lt;','p = 0.','p &lt; 0.001','***']))
print(f'{len(passed)} checks passed; 0 failed')
print('Includes cell-by-cell preservation checks for laboratory measurements and raw-library reconstruction of CBARP.')
