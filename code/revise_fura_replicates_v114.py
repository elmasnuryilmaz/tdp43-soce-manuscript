"""Correct the Fura-2 experimental unit after author confirmation (three wells on one plate)."""
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "manuscript/MANUSCRIPT_NEUROCHEMISTRY_INTERNATIONAL_FINAL.docx"
src = (ROOT / "code/condense_results_v107.py").read_text()
exec(src[src.index("def replace_span("):src.index("doc = Document(MANUSCRIPT)")])

doc = Document(PATH)
p = list(doc.paragraphs)

def change(idx, old, new):
    assert old in p[idx].text, (idx, old[:60])
    replace_span(p[idx], old, new)

change(0, "reduced store-operated Ca²⁺ entry", "a lower Ca²⁺-readdition response")
change(5,
       "Ca²⁺-readdition amplitude was 84% lower after knockdown (n = 3 independent cultures; Welch’s t-test p = 0.035). ER Ca²⁺ release was 33% lower (p = 0.31), and readdition normalised to each culture’s own release remained 77% lower. At 48 h, the WST-1 signal was 38.5% lower (n = 4).",
       "Ca²⁺-readdition amplitude was 84% lower in the knockdown wells of one culture plate (three wells per group; descriptive). ER Ca²⁺ release was 33% lower, and readdition normalised to each well’s own release was 77% lower. At 48 h, the WST-1 signal was 38.5% lower in four wells from one of three experiments (descriptive).")
change(5, "TDP-43 knockdown is associated with reduced SOCE", "This single-plate Fura-2 observation suggests a lower calcium-readdition response after TDP-43 knockdown")
change(54, "the samples of a group were prepared and measured on the same day.",
       "three samples per group came from separate wells of the same culture plate and were measured on the same day.")
change(54, "Each group comprised three independent cultures (n = 3), and both phases were read from the same recording of each culture.",
       "Each group comprised three wells on the same plate (n = 3 technical replicates), and both phases were read from the same recording of each well. No independent Fura-2 culture experiment is available for this comparison.")
change(59,
       "Fura-2 amplitudes from three independent cultures per group were compared with Welch’s two-tailed t-test (SciPy v1.12.0), which does not assume equal variances, and are reported with the difference in means and its 95% confidence interval; these two pre-specified comparisons were not adjusted for multiplicity. Effect sizes are Hedges’ g, the small-sample-corrected standardised mean difference, with the interval of Hedges and Olkin; at three samples per group the point estimate is large but its interval is correspondingly wide, so it describes the size of the difference rather than establishing it. With three samples per group the smallest two-sided p value attainable by a rank or permutation test is 0.10, so these comparisons rest on the parametric assumption of Welch’s test rather than on a distribution-free one. The readdition-to-release ratio of each culture was summarised descriptively. The 48-h WST-1 wells (n = 4) come from one experiment and were summarised descriptively, because wells within an experiment are not independent replicates.",
       "The Fura-2 data comprise three wells per group from one culture plate. We report their means, well-to-well SEM and observed differences descriptively; no inferential p value, confidence interval or standardised effect size is assigned to the Fura-2 comparison because wells from one plate are not independent biological replicates. The readdition-to-release ratio of each well was likewise summarised descriptively. The 48-h WST-1 values are four wells from one of three experiments and are also summarised descriptively; the other two experiments are not available in the package.")
change(66,
       "The readdition amplitude fell from 1.542 ± 0.282 to 0.245 ± 0.083 Δ(F340/F380) (mean ± SEM, n = 3 independent cultures per group; difference −1.30, 95% CI −2.40 to −0.19; Welch’s t-test p = 0.035; Hedges’ g = −2.9, 95% CI −5.3 to −0.5). ER Ca²⁺ release fell less and not significantly, from 0.268 ± 0.042 to 0.180 ± 0.061 (difference −0.09, 95% CI −0.30 to +0.13; p = 0.31; Hedges’ g = −0.8, 95% CI −2.2 to +0.6). Because both phases come from the same recording, readdition was also expressed relative to each culture’s own release. This ratio was 5.88 ± 1.10 in controls and 1.36 ± 0.01 after knockdown, 77% lower (Hedges’ g = −2.7, 95% CI −4.9 to −0.4), and every shTDP-43 culture had a lower ratio than every control culture (Supplementary Table S1). A smaller releasable store therefore does not account for most of the decrease.",
       "The readdition amplitude was 1.542 ± 0.282 in control wells and 0.245 ± 0.083 Δ(F340/F380) in knockdown wells (mean ± well-to-well SEM, three wells per group on one plate; difference −1.30, or 84% of the control mean). The observed ranges were 1.013–1.975 and 0.080–0.338, respectively. ER Ca²⁺ release was 0.268 ± 0.042 and 0.180 ± 0.061 (difference −0.088, or 33% of the control mean). Because both phases came from the same recording, readdition was also expressed relative to each well’s own release. This ratio was 5.88 ± 1.10 in controls and 1.36 ± 0.01 after knockdown, 77% lower; every knockdown well had a lower ratio than every control well (Supplementary Table S1). The ratio pattern suggests that a smaller releasable store alone may not explain the within-plate difference, but it cannot establish that this response recurs across independent cultures.")
change(69,
       "Panels C and D show the three independent cultures per group and mean ± SEM (readdition, Welch’s t-test p = 0.035; ER release, p = 0.31).",
       "Panels C and D show three wells per group from one culture plate and mean ± well-to-well SEM; the comparisons are descriptive and no inferential test is shown.")
change(132,
       "ER release fell by 33% and not significantly, whereas readdition fell by 84% and remained 77% lower when each culture was normalised to its own release, so a smaller releasable store does not account for most of the SOCE decrease.",
       "Within the single Fura-2 plate, mean ER release was 33% lower and mean readdition 84% lower; readdition remained 77% lower when each well was normalised to its own release. This pattern suggests that store release alone may not explain the smaller readdition response, but it needs replication in independent cultures.")
change(145,
       "The Fura-2 comparison rests on three independent cultures per group. At that size no rank or permutation test can reach a two-sided p below 0.10, so the difference is supported by a parametric test alone and would be more convincing with five to six cultures per group.",
       "The Fura-2 comparison rests on three wells per group from one culture plate, rather than independent biological experiments. Its 84% lower mean is a descriptive within-plate observation; the well-level Welch p value and confidence interval calculated previously cannot support a population-level inference. Independent culture experiments are required to estimate biological variability and test whether the response replicates.")
change(148,
       "TDP-43 depletion in SH-SY5Y cells was associated with a large reduction in Fura-2-measured SOCE and with altered expression and splicing of calcium-regulatory genes.",
       "In one SH-SY5Y culture-plate experiment, TDP-43 knockdown wells showed a lower Fura-2 calcium-readdition response. Independently, calcium-regulatory mRNA expression and splicing analyses identified candidate changes.")

doc.save(PATH)
print(PATH)
