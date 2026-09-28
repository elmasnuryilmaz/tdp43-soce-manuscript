"""List Supplementary Tables S18, S18b and S18c in the supplementary document."""
import copy
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SUPPLEMENTARY = ROOT / "supplementary/SUPPLEMENTARY_MATERIAL.docx"
ENTRIES = [
    "S18. Cell-composition adjustment of TRPC1 in the NYGC cohort: Cliff’s δ for ALS and for the "
    "comparison group against the same controls, unadjusted and after regression on SNAP25 or "
    "RBFOX3 with or without GFAP, together with the differences in the marker genes themselves.",
    "S18b. Cryptic STMN2 junction by group and region, on the samples of Table 5: samples and "
    "donors, samples with the junction detected, samples above 1% PSI, mean PSI, and the comparison "
    "group against controls. In cerebellum the junction was detected in three control samples and "
    "in no comparison-group sample; the resulting difference (δ = −0.08) is negligible.",
    "S18c. Within the comparison group, Spearman and partial (on SNAP25) correlations of TRPC1, "
    "SARAF, CBARP and SNAP25 with cryptic STMN2 PSI; and the Spearman correlation of TRPC1 with "
    "SNAP25 in every group and region.",
]
doc = Document(SUPPLEMENTARY)
s17 = next(p for p in doc.paragraphs if p.text.startswith("S17. "))
anchor = s17
for text in ENTRIES:
    new = copy.deepcopy(s17._p)
    for child in list(new):
        if child.tag.endswith("}r"):
            new.remove(child)
    anchor._p.addnext(new)
    from docx.text.paragraph import Paragraph
    anchor = Paragraph(new, s17._parent)
    anchor.add_run(text)
contents = next(p for p in doc.paragraphs if p.text.startswith("Supplementary data files\t"))
for r in contents.runs:
    if "S1–S17" in r.text:
        r.text = r.text.replace("S1–S17", "S1–S18")
assert "S1–S18" in contents.text
doc.save(SUPPLEMENTARY)
print("supplementary list updated")
