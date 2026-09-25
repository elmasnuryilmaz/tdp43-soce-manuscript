"""Apply the September 2026 reviewer-led editorial revision to the live DOCX.

Text edits operate on Word text nodes so embedded Zotero fields survive.
Run against the v1.0.3 manuscript only.
"""

from pathlib import Path
from docx import Document
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "manuscript/MANUSCRIPT_NEUROCHEMISTRY_INTERNATIONAL_FINAL.docx"


def replace(paragraph, old, new):
    nodes = paragraph._p.xpath(".//w:t")
    parts = [node.text or "" for node in nodes]
    whole = "".join(parts)
    if whole.count(old) != 1:
        raise ValueError(f"Expected exactly one occurrence of {old!r}; got {whole.count(old)}")
    start = whole.index(old)
    end = start + len(old)
    offsets = []
    position = 0
    for part in parts:
        offsets.append((position, position + len(part)))
        position += len(part)
    first = next(i for i, (a, b) in enumerate(offsets) if a <= start < b)
    last = next(i for i, (a, b) in enumerate(offsets) if a < end <= b)
    a, _ = offsets[first]
    z, _ = offsets[last]
    if first == last:
        nodes[first].text = parts[first][: start - a] + new + parts[first][end - a :]
    else:
        nodes[first].text = parts[first][: start - a] + new
        for i in range(first + 1, last):
            nodes[i].text = ""
        nodes[last].text = parts[last][end - z :]


def before(doc, anchor, paragraph):
    anchor._p.addprevious(paragraph._p)


def add_figure(doc, anchor, number, png, caption):
    doc.add_picture(str(png), width=Inches(6.15))
    picture = doc.paragraphs[-1]
    picture.paragraph_format.keep_with_next = True
    before(doc, anchor, picture)
    cap = doc.add_paragraph(style="Caption")
    cap.add_run(f"Figure {number}. ").bold = True
    cap.add_run(caption)
    before(doc, anchor, cap)


doc = Document(MANUSCRIPT)
p = list(doc.paragraphs)

replace(p[5], "while the Fura-2 Ca²⁺-readdition amplitudes averaged 1.542 and 0.245 Δ(F340/F380) in control and shTDP-43 cells, respectively (n = 3)", "while the Fura-2 Ca²⁺-readdition amplitude was approximately 84% lower after knockdown (n = 3)")
replace(p[5], "in six of seven regions", "in six brain regions")

replace(p[15], "; GENCODE v47", ". GENCODE v47")
replace(p[15], "; the transcript-level quantification", ". The transcript-level quantification")
replace(p[15], "; the junction and coverage analyses", ". The junction and coverage analyses")

replace(p[29], "Under a permissive definition (unannotated junction, ΔPSI ≥ 0.05, q < 0.05, control PSI ≤ 0.05, bootstrap lower bound above zero) the null produced almost as many calls as the real comparison in iPSC colonies (0.98 per real call) and twice as many in K562 total RNA (2.01; Supplementary Table S14), so raw call counts are not interpretable.", "The permissive definition required an unannotated junction, ΔPSI ≥ 0.05, q < 0.05, control PSI ≤ 0.05 and a bootstrap lower bound above zero (Supplementary Table S14).")
replace(p[34], " Excluding reference skips matters for an intronic index: counted as depth, the skips of spliced reads fill intronic windows, and in SH-SY5Y they would raise the STMN2 intron 2 index difference from +0.145 to +0.489.", "")

replace(p[46], "The junction-level analysis of Section 3.10", "The junction-level analysis of Section 3.9")
replace(p[46], "case numbers allowed this comparison in three regions.", "case numbers allowed this comparison in three regions. The public label does not identify the constituent disorders; these regions contained 49 cerebellar, 45 frontal-cortex and 35 temporal-cortex case samples (Table 5). Age, RNA integrity number and post-mortem interval were not included as covariates.")
replace(p[54], "GAPDH was the sole reference gene. In the four-target RNA set, its mean Ct was 19.49 in non-targeting shRNA controls and 19.50 in shTDP-43 cells (n = 4 each).", "GAPDH was used as the sole reference gene because its mean Ct was essentially unchanged in the four-target RNA set: 19.49 in non-targeting shRNA controls and 19.50 in shTDP-43 cells (n = 4 each).")
replace(p[61], "Fura-2 amplitudes (n = 3) and WST-1 well signals (n = 4) were summarised descriptively.", "Fura-2 amplitudes (n = 3) and 48-h WST-1 well signals (n = 4) were summarised descriptively, without inferential tests.")

replace(p[82], "A local-event method run on the same libraries gave a more conservative picture: SUPPA2 tested 31,474 events in 5,345 genes and returned no event that met the shared FDR and |ΔPSI| ≥ 0.10 thresholds, and among the 9,369 skipped-exon events whose coordinates matched between the two tools ΔPSI agreed only weakly (Spearman ρ = 0.319; 62.4% directional concordance).", "SUPPA2 gave a more conservative result on the same libraries. It tested 31,474 events in 5,345 genes and returned no event meeting the shared FDR and |ΔPSI| ≥ 0.10 thresholds. Among 9,369 skipped-exon events matched by coordinates, the tools agreed weakly on ΔPSI (Spearman ρ = 0.319; 62.4% directional concordance).")
replace(p[87], "identified a small set of events supported by adequate coverage and bootstrap intervals excluding zero (Table 3, Supplementary Figure S3).", "identified three events with adequate coverage and bootstrap intervals excluding zero: STIMATE, ORAI3 and STIM2 (Table 3, Supplementary Figure S3). A STIM1 event met the count and FDR criteria, but its interval included zero and is shown for context.")
replace(p[89], "The locus is therefore affected in every model except the iPSC-derived motor neurons", "The event-level evidence is shown in Figure 5. The locus is therefore affected in every model except the iPSC-derived motor neurons")

replace(p[95], "The null test also sets the limits of interpretation. It could be run in the three datasets with four control replicates:", "The null test also sets the limits of interpretation. Under the permissive rule, the null-to-real call ratios were 0.98 in iPSC colonies and 2.01 in K562 total RNA (Supplementary Table S14). The high-confidence null could be run in the three datasets with four control replicates:")
replace(p[96], "Under the permissive definition, SH-SY5Y additionally showed events in PLCD4 and in six genes of the expanded Ca²⁺ panel (ADCY1, SYT7, PDE3B, AKT3, PIK3CB, PPP2R5C), and novel combinations of annotated splice sites appeared in ATP2A3 (both doses, ΔPSI +0.16 and +0.20) and MCUB (iPSC colonies, +0.27) (Supplementary Figure S4B).", "Permissive-tier and annotated-site events are detailed in Supplementary Table S7 and Figure S4B.")
replace(p[97], "we report it as supported but not established.", "the two extraction methods therefore provide unequal support.")
replace(p[107], "In the primary SH-SY5Y model the positive control shifted only moderately:", "Excluding CIGAR reference skips affected this positive control: including them would have changed the STMN2 intron 2 index difference from +0.145 to +0.489. In the primary SH-SY5Y model, the skip-excluded index shifted only moderately:")

replace(p[140], "in six of seven ALS brain regions", "in six ALS brain regions")
replace(p[140], "TRPC1 remained unchanged in spinal cord and decreased in the neurological comparison cohorts.", "TRPC1 remained unchanged in the well-powered cervical and lumbar spinal-cord comparisons and decreased in the neurological comparison cohorts. Among the disease cohorts surveyed, ALS alone showed a significant TRPC1 increase, concentrated in brain regions.")
replace(p[138], "it was affected in five of six datasets across two species, had a median |ΔPSI| of 0.37 and a median of 246 junction reads per event, was independently supported by the annotation-free analysis and showed the largest expression change among the SOCE-regulatory genes.", "it was affected in five of six datasets across two species, was independently supported by the annotation-free analysis and showed the largest expression change among the SOCE-regulatory genes (Figure 5).")
replace(p[139], "but the pooled effect was negligible (ΔPSI +0.0013, 95% CI −0.022 to +0.024; p = 0.914).", "but the pooled effect was negligible (Supplementary Table S15).")
replace(p[141], "A coherent model is that", "The working model in Figure 6 is that")
replace(p[141], "depolarization-evoked", "depolarisation-evoked")

replace(p[250], "Table 3. Robust splicing events", "Table 3. Splicing events assessed for robustness")
replace(p[251], "Coordinates are 1-based and inclusive.", "Coordinates are 1-based and inclusive. The STIM1 interval crosses zero; that event is included for context and is not counted among the three interval-supported events.")
for paragraph in p:
    if "release v1.0.3" in paragraph.text:
        replace(paragraph, "v1.0.3", "v1.0.4")
        replace(paragraph, "v1.0.3", "v1.0.4")

add_figure(doc, p[91], 5, ROOT / "figures/main/Figure5_CBARP_splicing.png", "CBARP splicing across TDP-43-depletion datasets. (A) Genomic location and splice pattern of a representative iPSC-colony skipped-exon event, drawn from rMATS coordinates rather than read-coverage traces. (B) Per-library exon inclusion (PSI) for this event; each point represents one RNA-seq library. (C) Event-level ΔPSI for all 32 CBARP events meeting the prespecified FDR, effect-size and coverage filters in five datasets. Direction refers to the rMATS-defined form and varies by model; the evidence supports recurrent involvement of the locus, not a shared directional switch. Source: Supplementary Table S3.")

# Keep the two exploratory junction paragraphs within the preceding tissue section.
p[127]._element.getparent().remove(p[127]._element)

add_figure(doc, p[142], 6, ROOT / "figures/main/Figure6_working_model.png", "Working model of the observed calcium phenotype and candidate RNA changes. The laboratory Fura-2 assay found a smaller readdition amplitude after TARDBP knockdown in SH-SY5Y cells. Public knockdown RNA-seq identified changes in calcium-regulatory transcripts and recurrent CBARP splicing. ATP2A2, ORAI-family composition and SARAF are potential contributors to store filling, entry and feedback; CBARP is placed on a parallel voltage-gated channel branch. Arrows indicate hypotheses for future perturbation, not demonstrated causal links.")

lim_heading = doc.add_paragraph("Limitations", style="Heading 3")
before(doc, p[142], lim_heading)
lim1 = doc.add_paragraph("TARDBP depletion was verified by RT-qPCR but not at the protein level, and one TDP-43-targeting shRNA was used without a rescue experiment. The Fura-2 (n = 3) and 48-h WST-1 (n = 4) results are descriptive. The laboratory assay used undifferentiated SH-SY5Y cells, whereas the public SH-SY5Y RNA-seq study used an inducible model and the motor-neuron datasets did not include matched Ca²⁺ measurements. Cuvette Fura-2 ratios report net cytosolic accumulation and do not separate influx from extrusion or reuptake.")
before(doc, p[142], lim1)
lim2 = doc.add_paragraph("Bulk post-mortem expression can reflect cell composition, and age, RNA integrity and post-mortem interval were not modelled as covariates. The NMD dataset confounds TDP-43 status with sequencing batch, so its interaction results are exploratory.")
before(doc, p[142], lim2)
p[142]._element.getparent().remove(p[142]._element)

doc.save(MANUSCRIPT)
print(MANUSCRIPT)
